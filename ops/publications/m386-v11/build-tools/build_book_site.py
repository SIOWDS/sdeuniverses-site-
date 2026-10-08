from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import fitz,json,re,html,shutil,hashlib

ROOT=Path(__file__).resolve().parent
SITE=ROOT/'site';BOOK=SITE/'public/books/m/386';BOOK.mkdir(parents=True,exist_ok=True)
BASE='/books/m/386/';URL='https://sdeuniverses.com'+BASE
TITLE='隐私保护与“我”的诞生';SUB='AI时代的新经济典范';REV='20261008-m386-v11'
docx=ROOT/'output/隐私保护与我的诞生_第386卷_数字阅读版_v1.1.docx'
pdf=ROOT/'output/隐私保护与我的诞生_第386卷_数字阅读版_v1.1.pdf'
d=Document(docx);pd=fitz.open(pdf)
esc=html.escape
(BOOK/'downloads').mkdir(exist_ok=True)
shutil.copy2(docx,BOOK/'downloads/privacy-genesis-v1.1.docx')
shutil.copy2(pdf,BOOK/'downloads/privacy-genesis-v1.1.pdf')
pd[0].get_pixmap(matrix=fitz.Matrix(1.7,1.7)).save(BOOK/'cover.png')
from PIL import Image
Image.open(BOOK/'cover.png').convert('RGB').save(BOOK/'cover.jpg',quality=94)
(BOOK/'cover.png').unlink()
(BOOK/'vendor').mkdir(exist_ok=True)
for p in (SITE/'public/books/m/400/vendor').iterdir():
 if p.is_file():shutil.copy2(p,BOOK/'vendor'/p.name)

css='''*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#f8f6f0;color:#293342;font-family:Georgia,"Noto Serif SC","Songti SC",SimSun,serif;line-height:1.9}a{color:#1e2d50;text-underline-offset:4px}header,footer{padding:22px max(22px,calc((100vw - 1120px)/2));font:14px/1.6 sans-serif;border-bottom:1px solid #dcd6c7}header a{margin-right:15px}main{max-width:1120px;margin:auto;padding:54px 24px}h1,h2,h3{color:#1e2d50;line-height:1.5;font-weight:600}h1{font-size:clamp(28px,4.3vw,47px);margin:15px 0}h2{font-size:26px;margin-top:2.5em}h3{font-size:20px;margin-top:2em}.eyebrow{color:#9a742d;letter-spacing:.16em;font:13px/1.8 sans-serif}.hero{display:grid;grid-template-columns:310px 1fr;gap:65px;align-items:center}.cover{width:100%;height:auto;box-shadow:0 20px 42px #1e2d5025;border:1px solid #e1ddcf}.subtitle{font-size:24px;color:#927032}.meta{font:15px/2 sans-serif;color:#657080}.actions{display:flex;gap:12px;flex-wrap:wrap;margin:28px 0}.actions a{border:1px solid #d3c7ac;padding:10px 18px;text-decoration:none}.actions a.primary{background:#1e2d50;color:#fff;border-color:#1e2d50}.summary{font-size:18px;max-width:820px}.structure{display:grid;grid-template-columns:1fr 1fr;gap:0 50px}.toc-group{padding:20px 0;border-top:1px solid #dcd6c7}.toc-group h2{font-size:20px;margin:0 0 15px}.toc-group ol{padding-left:1.4em}.toc-group li{margin:10px 0}.reading{max-width:830px;font-size:20px;line-height:2.05}.reading p{text-align:justify;text-indent:2em;margin:.85em 0}.reading .note{text-indent:0;font-size:16px}.reading h1{font-size:32px}.reading .part{padding:60px 0 24px;border-top:1px solid #c9b07d}.reading h2{font-size:28px;margin-top:2.8em}.reading h3{font-size:21px;margin-top:2em}.reading article{scroll-margin-top:24px}.chapter-nav{display:flex;gap:16px;justify-content:space-between;font:15px/1.7 sans-serif;margin:50px 0;border-top:1px solid #dcd6c7;padding-top:20px}.tools{display:flex;gap:12px;font:15px/1.6 sans-serif;align-items:center}button{font:inherit;padding:8px 14px;background:white;border:1px solid #c9b07d;color:#1e2d50;cursor:pointer}.muted{color:#6b7482}.footnote{font:14px/1.9 sans-serif}.anchor{scroll-margin-top:20px}@media(max-width:720px){main{padding:30px 20px}.hero{grid-template-columns:1fr;gap:30px}.hero .cover{width:min(240px,65vw);margin:auto}.structure{grid-template-columns:1fr}.reading{font-size:18px}.reading h1{font-size:28px}.reading h2{font-size:25px}.actions a{padding:9px 13px}.meta{font-size:14px}}@media print{header,footer,.tools,.chapter-nav{display:none}body{background:white}.reading{max-width:none}}'''
(BOOK/'book.css').write_text(css)
def page(title,body,read=False,canonical=''):
 return f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(title)} · 德麦国际专著第386卷</title><meta name="description" content="{esc(TITLE+'：'+SUB+'。王德生、李佳城著。八编四十章，完整开放阅读。')}"><meta name="author" content="王德生、李佳城"><meta name="publication-revision" content="{REV}"><link rel="canonical" href="{URL+canonical}"><link rel="stylesheet" href="{BASE}book.css?v={REV}"></head><body><header><a href="/monographs/">德麦国际专著</a><a href="{BASE}">第386卷</a><a href="{BASE}chapters.html">全书目录</a><a href="{BASE}read.html">在线翻页</a></header><main class="{'reading' if read else ''}">{body}</main><footer>德麦国际出版社 · Demai International Press　｜　王德生、李佳城　｜　数字阅读版 v1.1</footer></body></html>'''
def markup(p):
 parts=[]
 for child in p._p:
  if child.tag==qn('w:hyperlink'):
   txt=''.join(child.itertext()) if False else ''.join(child.xpath('.//w:t/text()'))
   rid=child.get(qn('r:id'));href=d.part.rels[rid].target_ref if rid else ''
   parts.append(f'<a href="{esc(href,quote=True)}">{esc(txt)}</a>' if href.startswith('http') else esc(txt))
  elif child.tag==qn('w:r'):
   for c in child:
    if c.tag==qn('w:t'):parts.append(esc(c.text or ''))
    elif c.tag in [qn('w:br'),qn('w:cr')]:parts.append('<br>')
 return ''.join(parts)
units=[];parts=[];cur=None;body_seen=[]
for i,p in enumerate(d.paragraphs):
 if i<80:continue
 st=p.style.name;t=p.text
 if st=='Heading 1' and t!='全书目录':
  if cur:units.append(cur);cur=None
  parts.append({'title':t.replace('\n','　'),'index':len(units)})
  if t=='资料与方法说明':cur={'title':t,'slug':'sources','part':t,'paras':[]}
 elif st=='Heading 2':
  if cur:units.append(cur)
  n=len(units)
  slug='introduction' if n==0 else ('conclusion' if t.startswith('结语') else f'chapter-{n:02d}')
  cur={'title':t.replace('\n','　'),'slug':slug,'part':parts[-1]['title'] if parts else '导论','paras':[]}
 elif cur and st in ['Manuscript Body','Heading 3','Source Note']:
  tag='h3' if st=='Heading 3' else 'p';cl=' class="note"' if st=='Source Note' else ''
  cur['paras'].append(f'<{tag}{cl}>{markup(p)}</{tag}>')
  if st=='Manuscript Body':body_seen.append(t)
if cur:units.append(cur)
assert len(units)==43,len(units)
assert body_seen==[p.text for p in d.paragraphs if p.style.name=='Manuscript Body']
chapterdir=BOOK/'chapters';chapterdir.mkdir(exist_ok=True)
continuous=[];lastpart=None
for i,u in enumerate(units):
 if u['part']!=lastpart and u['part'] not in ['导论','资料与方法说明']:
  continuous.append(f'<div class="part"><div class="eyebrow">德麦国际专著 · 第386卷</div><h2>{esc(u["part"])}</h2></div>');lastpart=u['part']
 article=f'<article id="{u["slug"]}"><h2>{esc(u["title"])}</h2>'+''.join(u['paras'])+'</article>'
 continuous.append(article)
 nav='<nav class="chapter-nav">'
 if i:nav+=f'<a href="{units[i-1]["slug"]}.html">← 上一章</a>'
 nav+=f'<a href="{BASE}chapters.html">全书目录</a>'
 if i+1<len(units):nav+=f'<a href="{units[i+1]["slug"]}.html">下一章 →</a>'
 nav+='</nav>'
 (chapterdir/(u['slug']+'.html')).write_text(page(u['title'],f'<div class="eyebrow">{esc(u["part"])}</div>'+article+nav,True,'chapters/'+u['slug']+'.html'))
tocbody='';group=None
for u in units:
 if u['part']!=group:
  if group is not None:tocbody+='</ol></section>'
  group=u['part'];tocbody+=f'<section class="toc-group"><h2>{esc(group)}</h2><ol>'
 tocbody+=f'<li><a href="{BASE}chapters/{u["slug"]}.html">{esc(u["title"])}</a></li>'
tocbody+='</ol></section>'
(BOOK/'chapters.html').write_text(page('全书目录',f'<div class="eyebrow">CONTENTS · 386</div><h1>全书目录</h1><p>导论 · 八编四十章 · 结语 · 资料与方法说明</p><div class="structure">{tocbody}</div>',canonical='chapters.html'))
(BOOK/'text').mkdir(exist_ok=True)
front='<section><h2>阅读说明</h2>'+''.join(f'<p class="note">{markup(p)}</p>' for p in d.paragraphs[24:29])+'</section>'
(BOOK/'text/index.html').write_text(page(TITLE+' · 连续全文',f'<div class="eyebrow">第386卷 · 完整全文</div><h1>{TITLE}</h1><p class="subtitle">{SUB}</p><p class="note">王德生、李佳城 著　｜　正文200,195汉字</p>'+front+''.join(continuous),True,'text/'))
synopsis='隐私不仅关乎秘密资料的保管，也关乎一个尚在形成中的“我”能否继续发生。本书从个体性保护出发，讨论自我、现实与理念三界中的私人发生，人—AI复合主体、知识服务、我经济，以及共同规则如何承认“并非一切都必须成为我们的”。'
actions=f'<div class="actions"><a class="primary" href="read.html">友好阅读 · 在线翻页</a><a href="text/">连续全文</a><a href="chapters.html">分章阅读</a><a href="downloads/privacy-genesis-v1.1.pdf" download>下载全书 PDF</a><a href="downloads/privacy-genesis-v1.1.docx" download>下载 Word</a></div>'
hero=f'<section class="hero"><img class="cover" src="cover.jpg" alt="{esc(TITLE)}封面" width="915" height="1204"><div><div class="eyebrow">DEMAI INTERNATIONAL PRESS · 386</div><h1>{TITLE}</h1><p class="subtitle">{SUB}</p><p class="summary">我可以走向我们，而不必取消自己。</p><div class="meta">王德生、李佳城 著<br>ISBN 979-8-90690-877-3　·　US$20<br>2026年10月 · 数字阅读版 v1.1<br>八编四十章 · 正文200,195汉字 · 全书309页</div>{actions}</div></section>'
body=hero+f'<section class="summary"><h2>从保护秘密，到保护仍在发生的个体性</h2><p>{synopsis}</p><p>保护不自动带来创造，开放也不等于交出全部自我。书中把私人探索、对外服务、责任承担与持续修正分别展开，保留理论提案、构造案例、有限演示和经验研究之间的界限。</p></section><h2>全书结构</h2><div class="structure">{tocbody}</div><p class="footnote">本版完整开放网页、翻页与PDF阅读。书中案例与研究方案的证据状态见<a href="chapters/sources.html">资料与方法说明</a>。出版版式参照第220、302号专著。</p>'
(BOOK/'index.html').write_text(page(TITLE,body))

reader=(SITE/'public/books/m/220/read.html').read_text()
reader=reader.replace('我的三个宝贝',TITLE).replace('一个理论、一个工具、一声呼喊，与未来的主体经济学',SUB).replace('/books/m/220/','/books/m/386/')
cfg={'pdf':BASE+'downloads/privacy-genesis-v1.1.pdf','key':'books-m-386-v11','offset':6}
toc=[{'t':'封面','p':'','g':1,'l':1},{'t':'出版信息','p':'i','g':2,'l':1},{'t':'阅读说明','p':'ii','g':3,'l':1}]
for level,title,num in pd.get_toc():
 if level<3:toc.append({'t':title,'p':str(num-5) if num>=6 else 'iii','g':num,'l':level})
reader=re.sub(r'(<script type="application/json" id="cfg">).*?(</script>)',lambda m:m[1]+json.dumps(cfg,ensure_ascii=False)+m[2],reader,flags=re.S)
reader=re.sub(r'(<script type="application/json" id="toc">).*?(</script>)',lambda m:m[1]+json.dumps(toc,ensure_ascii=False)+m[2],reader,flags=re.S)
reader=reader.replace('https://cdnjs.cloudflare.com/ajax/libs/pdf.js/3.11.174/pdf.min.js',BASE+'vendor/pdf.min.js').replace('https://cdnjs.cloudflare.com/ajax/libs/pdf.js/3.11.174/pdf.worker.min.js',BASE+'vendor/pdf.worker.min.js')
reader=reader.replace('</head>',f'<meta name="publication-revision" content="{REV}"><meta name="author" content="王德生、李佳城"></head>')
reader=reader.replace('<script src="/taste/wds-companion/wds-read.js?v=20260817c" defer></script>','').replace('<script src="/wds-mode.js?v=20261005a" defer></script>','')
# Keep a useful fallback even if the script cannot load.
reader=reader.replace('<div id="msgT">正在载入书页…</div>',f'<div id="msgT">正在载入书页…</div><p><a style="color:#d9a441" href="{BASE}downloads/privacy-genesis-v1.1.pdf">打开全书 PDF</a> · <a style="color:#d9a441" href="{BASE}text/">网页全文</a></p>')
# Front matter must be directly reachable: make the visible page input use physical pages.
reader=reader.replace("$('pi').value = printed;", "$('pi').value = cur;")
reader=reader.replace("Math.max(1,total-CFG.offset+1)","total")
reader=reader.replace('第 <input id="pi"','书页 <input id="pi"')
reader=reader.replace('go((+this.value||1)+CFG.offset-1)', 'go(+this.value||1)')
(BOOK/'read.html').write_text(reader)

metadata={'id':'m-386','number':386,'title':TITLE,'subtitle':SUB,'authors':['王德生','李佳城'],'category':'culture','isbn':'9798906908773','price':'US$20','priceUsd':20,'currency':'USD','description':synopsis,'detailUrl':URL,'readUrl':URL+'read.html','flipUrl':URL+'read.html','textUrl':URL+'text/','chapterUrl':URL+'chapters.html','readMode':'full','openness':'full','readLabel':'友好阅读 · 在线翻页','pdfUrl':URL+'downloads/privacy-genesis-v1.1.pdf','docxUrl':URL+'downloads/privacy-genesis-v1.1.docx','coverUrl':URL+'cover.jpg','edition':'数字阅读版v1.1','publisher':'德麦国际出版社','pdfPages':len(pd),'hanTotal':200195,'revision':REV,'publishedAt':'2026-10-08T11:39:22Z','editionPublishedAt':'2026-10-08T11:39:22Z'}
(ROOT/'output/catalog-entry.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2))
manifest={'revision':REV,'title':TITLE,'subtitle':SUB,'authors':metadata['authors'],'isbn':'979-8-90690-877-3','price_usd':20,'pages':len(pd),'main_hanzi':200195,'body_paragraphs':len(body_seen),'chapter_units':42,'separate_reading_pages':43,'pdf_bookmarks':len(pd.get_toc()),'files':[]}
for p in sorted(BOOK.rglob('*')):
 if p.is_file() and p.name!='publication-manifest.json':manifest['files'].append({'path':str(p.relative_to(BOOK)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
(BOOK/'publication-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print('Book assets',len(manifest['files']),'units',len(units),'body preserved',len(body_seen))
