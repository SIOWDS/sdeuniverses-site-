from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import hashlib,json,re

root=Path(__file__).resolve().parent
src=root/'downloads/隐私保护与我的诞生_第386卷_全稿排版校样_v1.0.docx'
out=root/'output';out.mkdir(exist_ok=True)
d=Document(src)
body=[p.text for p in d.paragraphs if p.style.name=='Manuscript Body']
changes={
8:'数字阅读版 v1.1  ·  2026年10月',
20:'本次版本　数字阅读版 v1.1 · 2026年10月',
21:'本版以全稿排版校样v1.0为底本，参照第220、302号专著的190×250毫米开本、蓝金标题、宋体正文、分编页面与目录层级，完成数字阅读版整理。正文论述保持原稿，出版前置说明按本版用途更新。',
22:'本版供完整阅读、讨论与后续修订。书中理论提案、构造案例与候选研究方案应按各自的证据范围理解；数字出版不表示独立学术审读或完整经验验证已经完成。',
2168:'版本说明：本版依作者指令整理为数字阅读版v1.1，署名王德生、李佳城，德麦国际专著第386卷，ISBN 979-8-90690-877-3，定价US$20。正文保留原稿；完整原典对读、近邻理论系统比较与独立学术审读等后续工作，仍按书内资料说明所列范围推进。'
}
edits=[]
for i,t in changes.items():
 p=d.paragraphs[i];edits.append({'paragraph':i,'before':p.text,'after':t})
 # Preserve paragraph styling while replacing only publication statements.
 p.text=t
d.core_properties.title='隐私保护与“我”的诞生：AI时代的新经济典范'
d.core_properties.author='王德生、李佳城'
d.core_properties.subject='德麦国际专著第386卷 · 数字阅读版v1.1'
d.core_properties.keywords='隐私；主体性；SDE；我经济；人工智能'
target=out/'隐私保护与我的诞生_第386卷_数字阅读版_v1.1.docx'
d.save(target)
after=Document(target)
assert body==[p.text for p in after.paragraphs if p.style.name=='Manuscript Body']
report={'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'body_paragraphs':len(body),'body_hanzi':len(re.findall(r'[\u3400-\u4dbf\u4e00-\u9fff\U00020000-\U0002fa1f]',''.join(body))),'body_preserved':True,'publication_edits':edits}
(out/'editorial-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(target,report['body_paragraphs'],report['body_hanzi'])
