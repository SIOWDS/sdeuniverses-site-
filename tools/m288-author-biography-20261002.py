#!/usr/bin/env python3
"""A scoped author-biography replacement. No chapter or existing PDF page is reflowed."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from docx import Document
from docx.text.paragraph import Paragraph
from bs4 import BeautifulSoup
import fitz,hashlib,html,json,re,subprocess,sys
EXPECTED='5dd6f024e72b8a77661266f5e863a4ee4f0186c7718d39e0220c9e0a6ee2d720'
REV='m288-author-bio-20261002'
BIO=[
'陈凤，长期关注教育、女性成长、婚姻家庭、亲子关系，以及人工智能时代人的成长与教育。自2022年8月起跟随王德生博士持续学习，亲历其思想从“三视角”、321智慧，到SIO、SDE发生学不断生成与完善的过程。',
'几年的学习，也逐渐改变了我理解世界的方式：比起急着判断一个人、一件事“是什么”，我更愿意追问——它为什么会变成今天这样，又是怎样一步一步发生的？',
'我尝试把发生学带回真实生活，在孩子的成长、夫妻关系、女性处境、教育实践与自己的生命经验中不断观察、验证和修正理解。对我而言，发生学不是一套远离生活的理论，而是一种重新看见生活的方法：在学习中体验发生，在发生中理解自己、他人与世界。',
'现持续进行SDE发生学本体论下的阅读、研究。']
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
def sha(b):return hashlib.sha256(b).hexdigest()
def compact(s):return re.sub(r'\s+','',s)
def write_json(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def main():
    book,qa=map(lambda s:Path(s).resolve(),sys.argv[1:3]);qa.mkdir(parents=True,exist_ok=True)
    mp=book/'publication-manifest.json';m=json.loads(mp.read_text());old=deepcopy(m)
    assert m['release_sha256']==EXPECTED,'Newer release detected: stop, do not overwrite'
    assert (m['volume'],m['author'],m['isbn'],m['price_usd'],m['pdf_pages'])==(288,'陈凤','979-8-90690-433-1',20,383)
    for n,d in m['files'].items():assert sha((book/n).read_bytes())==d['sha256'],n
    write_json(qa/'previous-manifest.json',old);write_json(qa/'author-biography.json',BIO)
    docpath=book/'downloads/chen-feng-beauvoir-v1.0.docx'
    (qa/'baseline.docx').write_bytes(docpath.read_bytes())
    d=Document(docpath);paras=d.paragraphs
    start=next(i for i,p in enumerate(paras) if p.text=='作者介绍')
    end=next(i for i,p in enumerate(paras[start+1:],start+1) if p.text=='AI协作说明')
    before_text=[p.text for p in paras];oldbio=before_text[start+1:end]
    assert len(oldbio)==3 and oldbio[0].startswith('陈凤，SDE发生学的长期学习者')
    anchor=paras[start];template=deepcopy(paras[start+1]._p)
    for p in paras[start+1:end]:p._p.getparent().remove(p._p)
    for text in BIO:
        node=deepcopy(template);anchor._p.addnext(node);p=Paragraph(node,anchor._parent)
        for child in list(node):
            if child.tag!=W+'pPr':node.remove(child)
        p.add_run(text);anchor=p
    assert [p.text for p in d.paragraphs]==before_text[:start+1]+BIO+before_text[end:]
    d.save(docpath)
    with ZipFile(qa/'baseline.docx') as a,ZipFile(docpath) as b:
        assert a.namelist()==b.namelist()
        assert [n for n in a.namelist() if a.read(n)!=b.read(n)]==['word/document.xml']
    mini=Document(docpath);body=mini._element.body
    stop=next(p._p for p in mini.paragraphs if p.text.startswith('读者须知'))
    cut=False
    for node in list(body):
        if node is stop:cut=True
        if cut and node.tag!=W+'sectPr':body.remove(node)
    last=body.find(W+'sectPr')
    if last is not None:body.remove(last)
    body.append(deepcopy(d.sections[1]._sectPr))
    mini.save(qa/'author-front-matter.docx')
    profile=(qa/'lo-profile').as_uri()
    subprocess.run(['libreoffice','-env:UserInstallation='+profile,'--headless','--convert-to','pdf','--outdir',str(qa),str(qa/'author-front-matter.docx')],check=True,timeout=120)
    export=qa/'author-front-matter.pdf'
    pagepdf=fitz.open(export);assert len(pagepdf)==4
    assert compact(''.join(BIO)) in compact(pagepdf[3].get_text())
    assert 'AI协作说明' in compact(pagepdf[3].get_text())
    proof={'revision_id':REV,'author_bio_exact':True,'docx_only_biography_changed':True,'pdfs':[],'reindex_requested':False}
    for mode in ['print','reader']:
        path=book/f'downloads/chen-feng-beauvoir-{mode}-v1.0.pdf'
        original=fitz.open(path);updated=fitz.open(path);p=updated[3]
        assert len(original)==383 and compact(oldbio[0]) in compact(p.get_text())
        assert not p.get_links()
        newxref=updated.get_new_xref();updated.update_object(newxref,'<<>>');updated.update_stream(newxref,b'');p.set_contents(newxref)
        if mode=='reader':p.draw_rect(p.rect,color=None,fill=(251/255,248/255,240/255),overlay=False)
        p.show_pdf_page(p.rect,pagepdf,3,overlay=True)
        temp=qa/f'{mode}-updated.pdf';updated.save(temp,garbage=0,deflate=True);updated.close()
        check=fitz.open(temp)
        assert len(check)==383 and original.get_toc()==check.get_toc()
        changed=[]
        for i in range(383):
            if original[i].get_pixmap(matrix=fitz.Matrix(1,1)).samples!=check[i].get_pixmap(matrix=fitz.Matrix(1,1)).samples:changed.append(i+1)
            if i!=3:assert original[i].get_text()==check[i].get_text() and original[i].get_links()==check[i].get_links(),(mode,i)
        assert changed==[4],changed
        assert compact(''.join(BIO)) in compact(check[3].get_text())
        assert 'SDE发生学的长期学习者与写作者' not in compact(check[3].get_text())
        check[3].get_pixmap(matrix=fitz.Matrix(2,2)).save(qa/f'author-page-{mode}.png')
        proof['pdfs'].append({'type':mode,'pages':383,'changed_pages':changed,'other_382_pages_pixel_identical':True,'outline_and_links_preserved':True})
        check.close();original.close();path.write_bytes(temp.read_bytes());temp.unlink()
    textpath=book/'text/index.html';raw=textpath.read_text()
    soup=BeautifulSoup(raw,'html.parser')
    ledger=[(x.get('data-source-paragraph'),x.get_text()) for x in soup.select('[data-source-paragraph]')]
    section='<section id="author"><h2>作者介绍</h2>'+''.join('<p>'+html.escape(t)+'</p>' for t in BIO)+'</section>'
    raw,n=re.subn(r'<section id="author">.*?</section>',section,raw,count=1,flags=re.S);assert n==1
    textpath.write_text(raw)
    index=book/'index.html';raw=index.read_text();assert 'id="author"' not in raw
    needle='<section><h2>阅读与保存</h2>';assert raw.count(needle)==1
    index.write_text(raw.replace(needle,section+needle,1))
    mdpath=book/'downloads/chen-feng-beauvoir-v1.0.md';raw=mdpath.read_text()
    marker='## 内容提要\n';assert raw.count(marker)==1 and '## 作者介绍\n' not in raw
    bio_md='## 作者介绍\n\n'+'\n\n'.join(BIO)+'\n\n'
    changed_md=raw.replace(marker,bio_md+marker,1);assert changed_md.replace(bio_md,'',1)==raw
    mdpath.write_text(changed_md)
    for name in ['index.html','text/index.html','chapters.html','read.html']:
        p=book/name;raw=p.read_text()
        raw=re.sub(r'(chen-feng-beauvoir(?:-reader|-print)?-v1\.0\.(?:pdf|docx|md))(?=["\'])',r'\1?rev='+REV,raw)
        p.write_text(raw)
    check=BeautifulSoup(textpath.read_text(),'html.parser')
    assert ledger==[(x.get('data-source-paragraph'),x.get_text()) for x in check.select('[data-source-paragraph]')]
    assert len(check.select('section.chapter'))==68 and len(check.select('section.appendix'))==13
    for name in ['index.html','text/index.html']:
        s=BeautifulSoup((book/name).read_text(),'html.parser')
        assert [p.get_text() for p in s.select('#author > p')]==BIO
    ep=book/'edition.json';edition=json.loads(ep.read_text());edition['authorBiography']=BIO
    edition['revisionId']=REV;edition['authorBiographyUpdatedAt']='2026-10-02';edition['revisionNote']='仅按作者提供文字更新作者介绍；正文、页码和出版标识不变。';write_json(ep,edition)
    m['previous_release_sha256']=EXPECTED;m['revision_id']=REV;m['revision_description']='作者提供的正式简介；保留第一人称；正文与页码不变。'
    m['author_biography_sha256']=sha(''.join(BIO).encode());m['author_biography_updated_at']='2026-10-02'
    for n in m['files']:
        b=(book/n).read_bytes();m['files'][n]={'bytes':len(b),'sha256':sha(b)}
    m['release_sha256']=sha(json.dumps(m['files'],sort_keys=True,separators=(',',':')).encode());write_json(mp,m)
    proof['release_sha256']=m['release_sha256'];proof['body_html_paragraphs_unchanged']=len(ledger)
    proof['old_bio']=oldbio;proof['user_bio']=BIO
    proof['changed_book_files']=[n for n in m['files'] if m['files'][n]!=old['files'][n]]+['publication-manifest.json']
    write_json(qa/'author-bio-update-report.json',proof)
    print(json.dumps(proof,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
