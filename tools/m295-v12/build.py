#!/usr/bin/env python3
"""Build only volume 295 v1.2 from the authenticated v1.1 source.
No network requests, git writes, publication or indexing occurs in this builder.
"""
import argparse, copy, hashlib, html, json, re, shutil, subprocess, tempfile
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont
import fitz
HAN=re.compile(r'[\u4e00-\u9fff]')
def sha(b):return hashlib.sha256(b).hexdigest()
def han(t):return len(HAN.findall(t))
def norm(t):return re.sub(r'\s+','',t).replace('◆','')
def set_text(p,t):
    rpr=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for child in list(p._p):
        if child.tag not in (qn('w:pPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')):p._p.remove(child)
    run=p.add_run(t)
    if rpr is not None:run._r.insert(0,rpr)
def apply_string(s,replacements):
    for a,b in replacements.items():s=s.replace(a,b)
    return s

def new_table(doc,spec):
    n=len(spec['headers']);tb=doc.add_table(rows=1,cols=n);tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    width=130.0
    ratios=[.19,.81] if n==2 else ([.20,.43,.37] if spec['headers'][0]=='字段' else [1/3]*3)
    for col,ratio in zip(tb.columns,ratios):col.width=Mm(width*ratio)
    for rowvals in [spec['headers']]+spec['rows']:
        row=tb.rows[0] if rowvals is spec['headers'] else tb.add_row()
        trPr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trPr.append(nosplit)
        if rowvals is spec['headers']:trPr.append(OxmlElement('w:tblHeader'))
        for j,(cell,val) in enumerate(zip(row.cells,rowvals)):
            cell.width=Mm(width*ratios[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                k=OxmlElement('w:'+side);k.set(qn('w:w'),'75');k.set(qn('w:type'),'dxa');m.append(k)
            cp.append(m);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7ECF0' if rowvals is spec['headers'] else 'F9F8F4');cp.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=Pt(14);p.paragraph_format.keep_with_next=False
            run=p.add_run(val);run.font.name='Noto Serif CJK SC';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');run.font.size=Pt(9.3)
            if rowvals is spec['headers']:run.bold=True;run.font.color.rgb=RGBColor.from_string('23384A')
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'CFD6DB');borders.append(el)
    tb._tbl.tblPr.append(borders)
    return tb._tbl

def add_blocks(doc,anchor,blocks):
    elems=[]
    for block in blocks:
        if 'table' in block:elem=new_table(doc,block['table'])
        else:
            p=doc.add_paragraph(block['text'],style=block['style']);p.paragraph_format.widow_control=True
            if block['style']=='Heading 3':p.paragraph_format.keep_with_next=True
            elem=p._p
        anchor.addprevious(elem);elems.append((elem,block))
    return elems

def plates(source,dest):
    reg='/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc';sans='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    def font(n,serif=True):return ImageFont.truetype(reg if serif else sans,n,index=2)
    im=Image.open(source/'cover.jpg').convert('RGB');w,H=im.size;d=ImageDraw.Draw(im)
    d.rectangle((0,1175,w,H),fill='#071522');d.line((130,1200,w-130,1200),fill='#D7B678',width=2)
    def center(draw,t,y,f,fill):draw.text(((w-draw.textlength(t,font=f))/2,y),t,font=f,fill=fill)
    center(d,'存在主义已经有创造与揭示',1214,font(33),'#EDECE6');center(d,'中间态怎样记录',1260,font(33),'#EDECE6')
    center(d,'王德生 著',1323,font(27,False),'#EDECE6');center(d,'德麦国际出版社 · 第295卷 · 修订版 v1.2',1383,font(21,False),'#D7B678')
    im.save(dest/'cover.jpg',quality=95,subsampling=0);im.thumbnail((300,424));im.save(dest/'assets/cover-thumb.jpg',quality=92)
    im=Image.open(source/'backcover.jpg').convert('RGB');d=ImageDraw.Draw(im);d.rectangle((123,262,912,511),fill='#071522')
    for t,y in [('不把眼前的缺失，',275),('判成全部未来的终点；',322),('也不让“尚未”替成功担保。',369)]:d.text((153,y),t,font=font(34),fill='#EDECE6')
    d.line((154,429,350,429),fill='#D7B678',width=2);d.text((153,446),'让形成有记录，让失败能够结案。',font=font(26),fill='#EDECE6')
    im.save(dest/'backcover.jpg',quality=95,subsampling=0);im.thumbnail((230,325));im.save(dest/'assets/backcover-thumb.jpg',quality=92)

def patch_images(path,front,back):
    with ZipFile(path)as z:parts={n:z.read(n)for n in z.namelist()}
    parts['word/media/image3.jpg']=front.read_bytes();parts['word/media/image4.jpg']=back.read_bytes()
    with ZipFile(path,'w',ZIP_DEFLATED)as z:
        for n,b in parts.items():z.writestr(n,b)

def fragment(text):
    text=html.escape(text)
    return re.sub(r'\[(\d+)\]',lambda m:f'<a class="cite" href="/books/m/295/text/references/#ref{m[1]}">[{m[1]}]</a>',text)
def html_block(block,key):
    if 'table'in block:
        sp=block['table'];s='<div class="table-wrap"><table class="revision-table" data-source-paragraph="'+key+'"><thead><tr>'+''.join('<th>'+fragment(v)+'</th>'for v in sp['headers'])+'</tr></thead><tbody>'
        s+=''.join('<tr>'+''.join('<td>'+fragment(v)+'</td>'for v in row)+'</tr>'for row in sp['rows'])+'</tbody></table></div>'
    else:
        tag='h3' if block['style']=='Heading 3' else 'p';s=f'<{tag} data-source-paragraph="{key}">'+fragment(block['text'])+f'</{tag}>'
    return BeautifulSoup(s,'html.parser').contents[0]

def export_pdf(docx,out):
    with tempfile.TemporaryDirectory(prefix='m295-lo-')as tmp:
        cmd=['soffice','-env:UserInstallation='+Path(tmp,'profile').as_uri(),'--headless','--norestore','--convert-to','pdf:writer_pdf_Export','--outdir',str(out),str(docx)]
        res=subprocess.run(cmd,capture_output=True,text=True,timeout=180);print(res.stdout,res.stderr,flush=True)
        pdf=out/(docx.stem+'.pdf');assert res.returncode==0 and pdf.exists(),'PDF rendering failed'
    return pdf

def build(source,out,data_path):
    source=source.resolve();out=out.resolve();assert source!=out
    data=json.loads(data_path.read_text());origin=source/'downloads/xuwu-b295-v1.1.docx'
    assert sha(origin.read_bytes())==data['source_docx_sha256'],'Source docx changed; reconcile before editing'
    m=json.loads((source/'book.json').read_text());assert m['version']=='1.1' and m['number']==295 and m['priceUSD']==24
    if out.exists():shutil.rmtree(out)
    shutil.copytree(source,out);plates(source,out)
    doc=Document(origin);ps=list(doc.paragraphs);assert len(ps)==1998 and '第二十五章' in ps[1183].text
    replacements=data['global'];changes={int(k):v for k,v in data['replace'].items()}
    for part in [doc.part]+[p for p in doc.part.package.parts if p.partname.startswith('/word/header')]:
        if not hasattr(part,'element'):continue
        for node in part.element.xpath('.//w:t'):
            if node.text:node.text=apply_string(node.text,replacements).replace('数字阅读版 v1.1','数字阅读修订版 v1.2')
    for i,prefix in data.get('clause_prefix',{}).items():
        j=int(i)
        if j not in changes:changes[j]=prefix+ps[j].text
    for i,t in changes.items():set_text(ps[i],t)
    set_text(ps[14],'版本：2026年10月数字阅读修订版 v1.2');ps[17].paragraph_format.line_spacing=Pt(13.5)
    for r in ps[17].runs:r.font.size=Pt(9)
    additions=[]
    for i,blocks in data['insert_before'].items():additions.extend(add_blocks(doc,ps[int(i)]._p,blocks))
    for j in range(1844,1864):ps[j]._p.getparent().remove(ps[j]._p)
    additions.extend(add_blocks(doc,ps[1864]._p,data['appendix3']))
    count_styles={'Normal','EndLabel','EndQuote','EndNote'}
    retained=sum(han(p.text)for j,p in enumerate(ps)if 99<=j<1864 and not 1844<=j<1864 and p.style.name in count_styles)
    added=0
    for elem,b in additions:
        if 'table'in b:added+=sum(han(v)for row in [b['table']['headers']]+b['table']['rows']for v in row)
        elif b['style']in count_styles:added+=han(b['text'])
    total=retained+added;guide=sum(han(p.text)for p in ps[23:45]if p.style.name=='Normal')
    set_text(ps[15],f'开本：170 × 240毫米　｜　正文与附录：{total:,}汉字')
    doc.core_properties.title='虚无不是终点';doc.core_properties.subject=data['subtitle'];doc.core_properties.author='王德生';doc.core_properties.version='1.2';doc.core_properties.comments='根据应用创新审稿意见实质修订；外部活动与独立复填待验证。'
    doc.settings.element.append(OxmlElement('w:updateFields'));doc.settings.element[-1].set(qn('w:val'),'true')
    target=out/'downloads/xuwu-b295-v1.2.docx';doc.save(target);patch_images(target,out/'cover.jpg',out/'backcover.jpg')
    generated=export_pdf(target,out/'downloads');pdfpath=out/'downloads/xuwu-b295-reader-v1.2.pdf';generated.rename(pdfpath)
    pdf=fitz.open(pdfpath);pages=len(pdf);toc=pdf.get_toc();assert pages>299 and len(toc)>40
    webtoc=json.loads((source/'toc.json').read_text())
    for item in webtoc:
        item['title']=apply_string(item['title'],replacements);title=norm(item['title']);candidates=[x for x in toc if norm(x[1])==title]
        if not candidates and item['id']=='guide':candidates=[x for x in toc if x[1].startswith('导读')]
        if not candidates and item['kind']=='part':candidates=[x for x in toc if title in norm(x[1])]
        if candidates:item['pdf_page']=candidates[0][2]
        else:
            hits=[i+1 for i,p in enumerate(pdf)if title in norm(p.get_text()) and i>6];assert hits,('No actual PDF page for',item);item['pdf_page']=hits[0]
    assert len([e for e in webtoc if e['kind']=='chapter'])==35
    final=Document(target);page_map={}
    for i,p in enumerate(pdf):
        bottom=[b[4].strip() for b in p.get_text('blocks') if b[1]>p.rect.height-55];page_map[i+1]=next((v for v in bottom if re.fullmatch(r'[IVXLCDM]+|\d+',v)),'')
    wt={e['id']:e for e in webtoc}
    for p in final.paragraphs:
        if p.style.name not in {'TOC Entry','TOC Part'}:continue
        hl=p._p.xpath('./w:hyperlink');anchor=hl[0].get(qn('w:anchor'))if hl else '';key=anchor
        if anchor.startswith('ch_'):key='ch'+anchor.split('_')[-1].zfill(2)
        elif anchor.startswith('part_'):key='part'+anchor.split('_')[-1]
        elif anchor in {'app_1','app_2','app_3'}:key='appendix'+anchor[-1]
        if key in wt:
            label=page_map[wt[key]['pdf_page']]
            if label:
                nodes=p._p.xpath('./w:r/w:t')
                if nodes:nodes[-1].text=label
    final.save(target);pdf.close();generated=export_pdf(target,out/'downloads');generated.replace(pdfpath);pdf=fitz.open(pdfpath);assert len(pdf)==pages;assert len(pdf.get_toc())==len(toc)
    pdfmeta=pdf.metadata;pdfmeta.update(title='虚无不是终点',author='王德生',subject=data['subtitle'],keywords='SDE,存在主义,中间态,应用创新,第295卷,v1.2');pdf.set_metadata(pdfmeta);pdf.save(out/'downloads/tmp.pdf',garbage=4,deflate=True);pdf.close();(out/'downloads/tmp.pdf').replace(pdfpath)
    now='2026-10-02T00:00:00Z'
    m.update(subtitle=data['subtitle'],version='1.2',edition='2026年10月数字阅读修订版 v1.2',bodyHanzi=total,guideHanzi=guide,pdfPages=pages,description=data['description'],updatedAt=now,editionPublishedAt=now,evidenceStatus='文本修订完成；系列外真实活动、第二人独立复填及独立评分尚未完成',contentReviewDate='2026-10-02')
    for k in ['pdfUrl','docxUrl','coverUrl','backcoverUrl']:m[k]=m[k].replace('v1.1','v1.2').replace('v=1.1','v=1.2')
    (out/'book.json').write_text(json.dumps(m,ensure_ascii=False,indent=2)+'\n');(out/'toc.json').write_text(json.dumps(webtoc,ensure_ascii=False,indent=2)+'\n')
    replacements={**replacements,'数字阅读版 v1.1':'数字阅读修订版 v1.2','200,575':f'{total:,}','200575':str(total),'299页':f'{pages}页','299 页':f'{pages} 页','v=1.1':'v=1.2','v1.1.pdf':'v1.2.pdf','v1.1.docx':'v1.2.docx'}
    replacements[json.loads((source/'book.json').read_text())['description']]=data['description']
    for f in out.rglob('*.html'):
        s=BeautifulSoup(apply_string(f.read_text(),replacements),'html.parser')
        for i,t in changes.items():
            for e in s.select(f'[data-source-paragraph="{i}"]'):
                e.clear()
                for child in list(BeautifulSoup(fragment(t),'html.parser').contents):e.append(child)
        for i,blocks in data['insert_before'].items():
            e=s.select_one(f'[data-source-paragraph="{i}"]')
            if e:
                while e.parent and (e.parent.name=='blockquote' or e.parent.get('class')==['ledger']):e=e.parent
                for k,b in enumerate(blocks):e.insert_before(html_block(b,f'v12-{i}-{k}'))
            elif i=='1842':
                sec=s.select_one('#appendix2')
                if sec:
                    for k,b in enumerate(blocks):sec.append(html_block(b,f'v12-{i}-{k}'))
        sec=s.select_one('#appendix3')
        if sec:
            title=copy.copy(sec.select_one('h2'));sec.clear();sec.append(title);sec.append(BeautifulSoup('<div class="rule"></div>','html.parser').div)
            for k,b in enumerate(data['appendix3']):sec.append(html_block(b,f'v12-appendix3-{k}'))
        for script in s.select('script[type="application/ld+json"]'):
            try:
                obj=json.loads(script.string);obj.update(numberOfPages=pages,bookEdition=m['edition']);script.string=json.dumps(obj,ensure_ascii=False)
            except (ValueError,TypeError):pass
        for a in s.select('a[href*="read.html#page="]'):
            text=norm(a.get_text());matched=[e for e in webtoc if norm(e['title']) in text or text in norm(e['title'])and len(text)>6]
            if matched:a['href']=re.sub(r'#page=\d+',f"#page={matched[0]['pdf_page']}",a['href'])
        f.write_text(str(s))
    idx=out/'index.html';s=BeautifulSoup(idx.read_text(),'html.parser');section=s.new_tag('section',attrs={'class':'revision-note','id':'revision-v12'})
    section.append(BeautifulSoup('<h2>v1.2 实质修订说明</h2><p>副题与正文对齐；第三章写明四处原典对照的撤回与保留，第二十五章给出六字段记录，第三十三章加入加缪S→D→E路径切换，附录三提供独立复填与结案协议。</p><p><strong>证据边界：</strong>本版没有宣称已完成系列外真实活动验证或达到145分。记录模板不是已完成案例；修订版尚待独立复核。</p><p><a href="downloads/revision-note-v1.2.md">修订与证据状态说明</a> · <a href="text/appendix3/">六字段记录与复填协议</a></p>','html.parser'))
    s.select_one('section.hero').insert_after(section);idx.write_text(str(s));js=out/'assets/reader.mjs';js.write_text(js.read_text().replace('v1.1','v1.2').replace('v=1.1','v=1.2'))
    css=out/'assets/book.css';css.write_text(css.read_text()+'''\n/* v1.2: comparison/record tables, scoped to this volume. */
.table-wrap{overflow-x:auto;max-width:100%;margin:1.25em 0 1.5em}.revision-table{border-collapse:collapse;width:100%;min-width:540px;font-size:.90em;line-height:1.8}.revision-table th,.revision-table td{border:1px solid #cbd2d7;padding:.65em .8em;vertical-align:top;text-align:left}.revision-table th{background:#e7ecf0;color:#23384a}.dark .revision-table th{background:#263746;color:#eaeae4}.dark .revision-table td{border-color:#52606b}.revision-note{padding:1.2em 1.6em;border-left:3px solid #b89b65;background:rgba(180,160,120,.08);margin:2em 0}.hero h1{color:#e6e4de}@media(max-width:600px){.hero .sub{font-size:1.05rem;line-height:1.6}.revision-note{padding:1em}.revision-table{font-size:.9em}}
''')
    weights=[.2,.25,.2,.2,.15];score=lambda xs:round(sum(a*b for a,b in zip(xs,weights)),2)
    report={'volume':295,'version':'1.2','source_docx_sha256':data['source_docx_sha256'],'subtitle':data['subtitle'],'chapters':35,'parts':6,'appendices':3,'pdfPages':pages,'body_and_appendix_hanzi':total,'guide_hanzi':guide,'countConvention':'正文及附录的Normal、EndLabel、EndQuote、EndNote段落汉字，加新增表格单元格汉字；不含标题、章首引句、导读、参考文献、出版信息。与旧版200575口径衔接。','replaced_source_paragraphs':len(changes),'inserted_blocks':len(additions),'old_review_score':score([142,134,132,128,140]),'review_target_as_written':145.10,'review_design_target_recalculated':score([148,146,144,142,146]),'new_edition_independently_scored':False,'external_case_completed':False,'independent_refill_completed':False,'real_preregistered_closure_completed':False,'reindex':False,'source_revision_sha256':sha(data_path.read_bytes())}
    full=BeautifulSoup((out/'text/index.html').read_text(),'html.parser');report['html_source_blocks']=len(full.select('[data-source-paragraph]'));assert len(full.select('section.chapter'))==35
    expected={e['data-source-paragraph']:norm(e.get_text())for e in full.select('[data-source-paragraph]')}
    for k,t in changes.items():
        if str(k)in expected:assert expected[str(k)]==norm(t),(k,'revised source mismatch')
    for f in (out/'text').glob('*/index.html'):
        s=BeautifulSoup(f.read_text(),'html.parser')
        for e in s.select('[data-source-paragraph]'):assert expected[e['data-source-paragraph']]==norm(e.get_text()),str(f)
    report['fulltext_chapters_equal']=True;report['docx_sha256']=sha(target.read_bytes());report['pdf_sha256']=sha(pdfpath.read_bytes())
    note=f'''# 《虚无不是终点》第295卷：v1.2实质修订记录

2026年10月2日。王德生著；德麦国际出版社；ISBN 979-8-90690-362-4；US$24。

新副题：{data['subtitle']}。

## 已完成的文本修订

保留六部35章、三份附录及30项原有参考资料。同步修订导读、前言、第二章、第三章、第四章、第六章、第七章、第五部标题、第二十五章、第三十三章、结语与附录二、三。不是以附加免责声明替代正文修改。

第三章新增四处对照，明确撤回全称遗漏；第二十五章把三张白纸转为六字段记录；第三十三章增加加缪S→D→E的有条件路径切换，并列明误读代价与失败结果；结语写明微选择及副题的实际回写；附录二采用A2-01至A2-09条款号；附录三改为真实活动、独立复填、近邻重述及结案的空白协议。封面副题和封底判断相应调整，沿用原图主体。

## 没有声称完成的工作

没有编造系列外真实活动，没有代填第二人的独立结论，没有把文本撤回冒充预注册活动结案。v1.1独立应用阅读分134.90经脚本重算一致。原审稿意见把目标总分写为145.10；按原列五维及权重重算为145.20，已明确标出算术差异。v1.2尚未独立评分，不标为认证达标。

## 规模与校验

PDF {pages}物理页；正文与附录{total:,}汉字，导读另{guide:,}汉字。字数口径详见revision-report-v1.2.json；不把全文件或标题统计冒充正文。全书35章，网页与逐章正文逐块一致。发布日期沿用原首发日期；本次作为更新，不另算一本新书。旧版下载文件保留，新入口指向v1.2。

生成器仅处理本书，不执行reindex。在线发布与验证状态另以实际部署验收为准。
'''
    (out/'downloads/revision-note-v1.2.md').write_text(note);(out/'downloads/revision-report-v1.2.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    manifest={'volume':295,'version':'1.2','source_revision_sha256':report['source_revision_sha256'],'pdfPages':pages,'bodyHanzi':total,'html_source_blocks':report['html_source_blocks'],'files':{}}
    for f in out.rglob('*'):
        if not f.is_file()or f.name=='publication-manifest.json':continue
        rel=str(f.relative_to(out))
        if rel in ['downloads/xuwu-b295-reader-v1.1.pdf','downloads/xuwu-b295-v1.1.docx']:continue
        assert f.suffix.lower()not in ['.ttf','.otf','.ttc','.woff','.woff2'];manifest['files'][rel]={'size':f.stat().st_size,'sha256':sha(f.read_bytes())}
    (out/'publication-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,ensure_ascii=False,indent=2),flush=True)
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('source',type=Path);ap.add_argument('output',type=Path);ap.add_argument('revision',type=Path);a=ap.parse_args();build(a.source,a.output,a.revision)
