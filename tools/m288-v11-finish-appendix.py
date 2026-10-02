#!/usr/bin/env python3
"""Finish the reviewed v1.1 book: one appendix orphan, real TOC, metadata, exact old-URL aliases."""
from pathlib import Path
from collections import defaultdict
import sys,os,shutil,json,hashlib,tempfile,subprocess,re,unicodedata
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from bs4 import BeautifulSoup
import fitz

def norm(t):return re.sub(r'\s+','',unicodedata.normalize('NFKC',t))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def js(p,v):Path(p).write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n')
def convert(path,out):
    out.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='m288-final-') as profile:
        r=subprocess.run(['libreoffice','-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf:writer_pdf_Export','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=180)
        pdf=out/(path.stem+'.pdf');assert r.returncode==0 and pdf.exists(),r.stdout+r.stderr
    return pdf

def main():
    incoming,out=map(Path,sys.argv[1:3]);manifests=list(incoming.rglob('book/publication-manifest.json'));assert len(manifests)==1
    source=manifests[0].parent;m=json.loads(manifests[0].read_text())
    assert m['release_sha256']=='56d3226b275c071c7604e2db501e49e17821bc4ff0ff15ba44339a4361da9f20'
    assert (m['volume'],m['isbn'],m['price_usd'],m['pdf_pages'])==(288,'979-8-90690-433-1',20,354)
    for n,v in m['files'].items():assert sha(source/n)==v['sha256'],n
    book=out/'book';qa=out/'qa';qa.mkdir(parents=True,exist_ok=True);shutil.copytree(source,book,dirs_exist_ok=True)
    path=book/'downloads/chen-feng-beauvoir-v1.1.docx';d=Document(path);before=[p.text for p in d.paragraphs];inside=False;modified=0
    for p in d.paragraphs:
        b=p._p.findall(qn('w:bookmarkStart'));a=b[0].get(qn('w:name')) if b else ''
        if a=='appendix09':inside=True
        elif a=='appendix10':inside=False
        if not inside:continue
        f=p.paragraph_format
        if p.style.name=='Normal':f.line_spacing=Pt(17.4);f.space_after=Pt(2.7);modified+=1
        elif p.style.name=='Heading 3':f.space_before=Pt(7);f.space_after=Pt(4);modified+=1
        if p.text.startswith('目前最可辩护的定位仍然是：'):f.keep_together=True
    assert modified==27 and [p.text for p in d.paragraphs]==before
    d.save(path)
    for attempt in range(12):
        pdfpath=convert(path,out/f'conversion-{attempt}');p=fitz.open(pdfpath);by=defaultdict(list)
        for lev,t,pg in p.get_toc():by[norm(t)].append(pg)
        mapping={};names={}
        for para in d.paragraphs:
            if para.style.name not in ['Heading 1','Heading 2']:continue
            b=para._p.findall(qn('w:bookmarkStart'))
            if not b:continue
            a=b[0].get(qn('w:name'));assert by.get(norm(para.text)),para.text
            mapping[a]=by[norm(para.text)][0];names[a]=para.text
        changes=0
        for para in d.paragraphs:
            if para.style.name not in ['BookTOC','BookTOCPart']:continue
            links=para._p.findall(qn('w:hyperlink'))
            if not links:continue
            link=links[0];a=link.get(qn('w:anchor'));ts=link.findall('.//'+qn('w:t'))
            if ts:
                changes+=ts[0].text!=names[a];ts[0].text=names[a]
                for t in ts[1:]:t.text=''
            ts=para._p.findall(qn('w:r'))[-1].findall(qn('w:t'))
            if ts:changes+=ts[-1].text!=str(mapping[a]);ts[-1].text=str(mapping[a])
        print('Contents convergence',attempt,len(p),changes,flush=True);p.close();d.save(path)
        if not changes:break
    else:raise RuntimeError('TOC did not converge; no publication permitted')
    pdf=fitz.open(pdfpath);pages=len(pdf);assert pages==353,pages
    toc=sorted([[2 if a.startswith('ch') else 1,names[a],n] for a,n in mapping.items()],key=lambda v:v[2])
    toc=[[1,'封面',1],[1,'书名页',2]]+toc+[[1,'封底',pages]]
    prev=0;clean=[]
    for lev,t,pg in toc:lev=min(lev,prev+1);clean.append([lev,t,pg]);prev=lev
    pdf.set_toc(clean);meta=pdf.metadata;meta.update({'title':'陈凤讲波伏娃：普通人都能懂','author':'陈凤','subject':'德麦国际专著第288卷 · 数字阅读版v1.1 · 逐页精修版','keywords':'波伏娃,陈凤,普通人都能懂,SDE,979-8-90690-433-1'});pdf.set_metadata(meta)
    pp=book/'downloads/chen-feng-beauvoir-print-v1.1.pdf';pp.unlink();pdf.save(pp,garbage=4,deflate=True);pdf.close()
    pdf=fitz.open(pp)
    for i,page in enumerate(pdf):
        if i not in [0,pages-1]:page.draw_rect(page.rect,color=None,fill=(0.988,0.984,0.965),overlay=False)
    rp=book/'downloads/chen-feng-beauvoir-reader-v1.1.pdf';rp.unlink();pdf.save(rp,garbage=4,deflate=True);pdf.close()
    for f in book.rglob('*.html'):
        soup=BeautifulSoup(f.read_text(),'html.parser')
        for tag in soup.select('meta[name="description"]'):tag['content']=tag.get('content','').replace('383页',str(pages)+'页').replace('354页',str(pages)+'页')
        for tag in soup.select('.statgrid strong'):
            if tag.get_text() in ['383','354']:tag.string=str(pages)
        for tag in soup.select('script[type="application/ld+json"]'):
            v=json.loads(tag.string);v['numberOfPages']=pages;tag.string=json.dumps(v,ensure_ascii=False)
        if f.name=='read.html':soup.find(id='toc').string=json.dumps([{'t':t,'p':str(pg),'g':pg,'l':lev} for lev,t,pg in clean],ensure_ascii=False)
        f.write_text(str(soup))
    edition=json.loads((book/'edition.json').read_text());edition['pdfPages']=pages;js(book/'edition.json',edition)
    for suffix,ext in [('','docx'),('-print','pdf'),('-reader','pdf'),('','md')]:shutil.copy2(book/f'downloads/chen-feng-beauvoir{suffix}-v1.1.{ext}',book/f'downloads/chen-feng-beauvoir{suffix}-v1.0.{ext}')
    p=fitz.open(pp);r=fitz.open(rp);assert all(norm(a.get_text())==norm(b.get_text()) for a,b in zip(p,r))
    texts=[''.join(pg.get_text().split()) for pg in p];geometry=[];audit=[]
    for i,pg in enumerate(p):
        spans=[sp for bl in pg.get_text('dict')['blocks'] for ln in bl.get('lines',[]) for sp in ln['spans']]
        geometry.append([[sp['text'],[round(x,2) for x in sp['bbox']],round(sp['size'],2),sp['font']] for sp in spans])
        body=[sp for sp in spans if sp['text'].strip() and 49<sp['bbox'][1] and sp['bbox'][3]<pg.rect.height-32]
        overflow=[sp['text'] for sp in body if sp['bbox'][0]<35 or sp['bbox'][2]>pg.rect.width-25]
        audit.append({'page':i+1,'overflow':overflow,'replacement_glyph':'\ufffd' in pg.get_text(),'body_characters':sum(len(norm(x['text'])) for x in body),'pdf_pair_text_equal':True})
    tx=hashlib.sha256(json.dumps(texts,ensure_ascii=False,separators=(',',':')).encode()).hexdigest();ge=hashlib.sha256(json.dumps(geometry,ensure_ascii=False,separators=(',',':')).encode()).hexdigest()
    js(qa/'page-map.json',mapping);js(qa/'heading-names.json',names);js(qa/'final-page-audit.json',audit);js(qa/'all-page-geometry.json',geometry)
    js(qa/'render-signatures.json',{'pages':pages,'text_sha256':tx,'geometry_sha256':ge,'expected_text_sha256':'277a621f90ef9851958c338894279a89765e4bfff5f45e36a5be10e8dd7c4f51','expected_geometry_sha256':'43d7adaf7fa20a5e4c42e9e07fd5dcd5441750f07033461de9148bb48b58fa5e'})
    assert not any(x['overflow'] or x['replacement_glyph'] for x in audit)
    assert tx=='277a621f90ef9851958c338894279a89765e4bfff5f45e36a5be10e8dd7c4f51','Rendered page text differs from visually checked 353-page edition'
    assert ge=='43d7adaf7fa20a5e4c42e9e07fd5dcd5441750f07033461de9148bb48b58fa5e','Rendered page geometry differs from visually checked 353-page edition'
    active={f.relative_to(book).as_posix():{'bytes':f.stat().st_size,'sha256':sha(f)} for f in sorted(book.rglob('*')) if f.is_file() and f.name!='publication-manifest.json'}
    assert not any(Path(n).suffix.lower() in ['.ttf','.otf','.ttc','.woff','.woff2','.env'] for n in active)
    m.update({'pdf_pages':pages,'files':active,'release_sha256':hashlib.sha256(json.dumps(active,sort_keys=True,separators=(',',':')).encode()).hexdigest()});js(book/'publication-manifest.json',m)
    js(qa/'appendix-fit.json',{'modified_paragraphs':modified,'body_font_pt':11,'body_line_pt':17.4,'same_text_and_geometry_as_reviewed_edition':True})
    print('READY',pages,m['release_sha256'])
if __name__=='__main__':main()
