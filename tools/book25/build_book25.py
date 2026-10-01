#!/usr/bin/env python3
"""Build only the authorized book-25 publication. No indexing or deployment calls."""
from __future__ import annotations
import argparse, hashlib, html, json, math, re, shutil, urllib.request
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from weasyprint import HTML
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fitz

TITLE='心血管健康的发生学重建'
SUB='从风险控制到调节能力与生活重建'
AUTHORS='王德生、秦莉'
ISBN='978-1-970820-33-1'
BASE='/books/m/25/'
VERSION='20261002-v1'
NAVY='#172b49'; GOLD='#b28c43'; PAPER='#fbf8f0'
PDFNAME='book25-reader-v1.pdf'; DOCNAME='book25-editable-v1.docx'

def write(path:Path, text:str):
    path.parent.mkdir(parents=True,exist_ok=True);path.write_text(text,encoding='utf-8')
def digest(path:Path):return hashlib.sha256(path.read_bytes()).hexdigest()
def han(s:str):return len(re.findall(r'[\u3400-\u4dbf\u4e00-\u9fff]',s))
def escaped(s:str):return html.escape(s,quote=True)
def inline(s:str, prefix:str=''):
    t=escaped(s)
    t=re.sub(r'https?://[^\s<>；。]+',lambda m:'<a class="url" href="'+m[0]+'" target="_blank" rel="noopener">'+m[0]+'</a>',t)
    t=re.sub(r'(?<![\w/])\[(\d+)(?:[—–-](\d+))?\]',lambda m:'<a class="cite" href="'+prefix+'#ref-'+m[1]+'">'+m[0]+'</a>',t)
    return t

def parse(files:list[Path]):
    sections=[]; current=None; source=[]
    for f in files:
        text=f.read_text(encoding='utf-8').replace('进入第二、第二十八编章相关讨论，即','进入第七、第八、第十及第二十八章等讨论，即')
        text=text.replace('进入第七、第八、第十及第二十八章等讨论，即第七、第八、第十及第二十八章','进入第七、第八、第十及第二十八章')
        source.append(text.strip())
        for block in re.split(r'\n\s*\n',text.strip()):
            block=block.strip()
            if not block:continue
            m=re.match(r'^(#{1,3})\s+(.+)$',block,re.S)
            if m:
                level=len(m[1]);title=m[2].strip()
                chapter=level==2 and bool(re.match(r'^第.{1,4}章',title))
                if level==1 or chapter:
                    kind='chapter' if chapter else ('part' if re.match(r'^第.{1,3}编',title) else ('appendix' if title.startswith('附录') else ('references' if title.startswith('参考文献') else 'front')))
                    current={'id':f's{len(sections)+1:03d}','title':title,'kind':kind,'blocks':[],'source':f.name}
                    sections.append(current)
                else:
                    assert current is not None
                    current['blocks'].append({'type':'h3','text':title})
            else:
                assert current is not None
                typ='quote' if block.startswith('> ') else 'p'
                value=block[2:].strip() if typ=='quote' else block.replace('\n',' ')
                if typ=='quote' and value.startswith('带走的话：'):typ='takeaway'
                current['blocks'].append({'type':typ,'text':value})
    assert sum(s['kind']=='chapter' for s in sections)==30
    assert sum(s['kind']=='part' for s in sections)==6
    assert sum(s['kind']=='appendix' for s in sections)==5
    refs=[b for s in sections if s['kind']=='references' for b in s['blocks'] if re.match(r'^\[\d+\]',b['text'])]
    assert len(refs)==40
    return sections,'\n\n'.join(source)+'\n'

def section_html(s:dict, prefix:str=''):
    level='h2' if s['kind']=='chapter' else 'h1'
    label={'chapter':'CHAPTER','part':'PART','appendix':'APPENDIX','references':'REFERENCES','front':'DEMAI MONOGRAPHS · 025'}[s['kind']]
    out=[f'<section class="unit {s["kind"]}" id="{s["id"]}"><div class="kicker">{label}</div><{level}>{escaped(s["title"])}</{level}><div class="gold-rule"></div>']
    for b in s['blocks']:
        text=inline(b['text'],prefix);typ=b['type']
        if typ=='h3':out.append('<h3>'+text+'</h3>')
        elif typ in ('quote','takeaway'):out.append(f'<blockquote class="{typ}"><p>{text}</p></blockquote>')
        else:
            match=re.match(r'^\[(\d+)\]',b['text']); rid=f' id="ref-{match[1]}"' if s['kind']=='references' and match else ''
            cls='ref' if s['kind']=='references' else ('formula' if 'S=F(D,E)' in b['text'] and len(b['text'])<150 else '')
            out.append(f'<p class="{cls}"{rid}>'+text+'</p>')
    out.append('</section>');return '\n'.join(out)

PRINT_CSS='''
@page{size:190mm 250mm;margin:19mm 20mm 18mm;background:#fffefa;
 @top-right{content:string(running);font-family:'Noto Sans CJK SC',sans-serif;font-size:7pt;color:#888274}
 @bottom-center{content:counter(page);font-family:'Noto Sans CJK SC',sans-serif;font-size:8pt;color:#8b806a}
}
@page cover{margin:0;background:#172b49;@top-right{content:none}@bottom-center{content:none}}
*{box-sizing:border-box}body{font-family:'Noto Serif CJK SC',serif;font-size:11.2pt;line-height:1.85;color:#2c2d30;margin:0}
a{color:inherit;text-decoration:none}p{margin:0 0 3.2mm;text-indent:2em;text-align:justify;orphans:3;widows:3}
.cover{page:cover;width:190mm;height:250mm;break-before:page;break-after:page}.cover img{display:block;width:190mm;height:250mm}
.unit{break-before:page}.unit h1,.unit h2{font-family:'Noto Serif CJK SC',serif;color:#172b49;line-height:1.5;margin:4mm 0 3mm;bookmark-level:1;font-weight:700;string-set:running content()}
.unit h1{font-size:21pt}.unit h2{font-size:19pt;bookmark-level:2}
.kicker{font-family:'Noto Sans CJK SC',sans-serif;font-size:8pt;letter-spacing:.12em;color:#b28c43;padding-top:3mm}
.gold-rule{height:.6mm;width:16mm;background:#b28c43;margin:5mm 0 8mm}
h3{font-size:12.2pt;color:#172b49;margin:6mm 0 3mm;line-height:1.6;break-after:avoid;bookmark-level:none}
blockquote{margin:4mm 0 7mm;padding:3mm 4mm;border-left:.8mm solid #b28c43;background:#f5f0e4;break-inside:avoid}
blockquote p{text-indent:0;text-align:left;margin:0;color:#4d5260;font-size:10.8pt;line-height:1.8}
.takeaway{border-left:0;border-top:.4mm solid #b28c43;border-bottom:.4mm solid #b28c43;background:none;margin-top:8mm;padding:4mm 0}.takeaway p{color:#172b49;font-weight:700}
.part{break-after:page;padding-top:40mm}.part h1{font-size:27pt}.part p{text-indent:0;color:#6f6a60;font-size:12pt;margin-top:16mm}
.references p{font-family:'Noto Serif CJK SC',serif;font-size:9pt;line-height:1.65;text-indent:0;text-align:left;margin-bottom:4mm;overflow-wrap:anywhere;word-break:normal}
.url{overflow-wrap:anywhere;word-break:break-all}.cite{font-size:.82em;color:#6d7683}.formula{text-indent:0;text-align:center}
.toc{break-before:page;break-after:page}.toc h1{font-size:22pt;color:#172b49;bookmark-level:none}.toc ul{list-style:none;padding:0}.toc li{margin:1.8mm 0;font-size:10pt;line-height:1.55;break-inside:avoid}.toc li.c{padding-left:5mm}.toc li.p{margin-top:4mm;font-weight:bold;color:#172b49}.toc a::after{content:leader('.') target-counter(attr(href),page)}
'''

WEB_CSS='''
:root{--paper:#fbf8f0;--ink:#2d2d2d;--navy:#1f3a5f;--gold:#b08a3c;--dim:#8a8578;--line:#e4dccb;--soft:#f3eee1}
html.dark{--paper:#121518;--ink:#e3dfd6;--navy:#9fc0e6;--gold:#d9b45c;--dim:#9caaa8;--line:#2a3036;--soft:#1a1f24}
*{box-sizing:border-box}html{scroll-padding-top:80px}body{margin:0;background:var(--paper);color:var(--ink);font:18px/2 'Noto Serif CJK SC','Songti SC','SimSun',serif}
a{color:var(--navy);text-underline-offset:.2em}.bar{position:sticky;top:0;z-index:5;display:flex;flex-wrap:wrap;gap:.7rem;align-items:center;padding:.6rem 1rem;background:var(--paper);border-bottom:1px solid var(--line);font:14px/1.5 sans-serif}.bar a{text-decoration:none}.bar .sp{margin-left:auto}.bar button{font:inherit;background:none;border:1px solid var(--line);border-radius:5px;padding:.2rem .55rem;color:var(--ink);cursor:pointer}
#prog{position:fixed;top:0;left:0;height:3px;background:var(--gold);width:0;z-index:8}.wrap{max-width:780px;margin:auto;padding:1.5rem 1.3rem 5rem}.hero{text-align:center;margin:1.7rem 0 3rem}.hero img{width:210px;box-shadow:0 18px 40px #0003}.hero h1{font-size:1.9rem;color:var(--navy);line-height:1.6}.hero p{text-indent:0;text-align:center;font-size:.92rem;color:var(--gold)}
p{margin:0 0 .95em;text-indent:2em;text-align:justify;overflow-wrap:break-word}.kicker{font:12px/1.5 sans-serif;color:var(--gold);letter-spacing:.15em;margin-top:1rem}.unit{padding-top:1rem;margin:3.5rem 0}.unit h1,.unit h2{font-size:1.55rem;color:var(--navy);line-height:1.6;margin:.7rem 0}.gold-rule{height:2px;width:46px;background:var(--gold);margin:1rem 0 1.6rem}h3{font-size:1.12rem;color:var(--navy);line-height:1.7;margin:2.2rem 0 .7rem}.part{border-top:1px solid var(--line);padding-top:3rem;margin-top:5rem}.part h1{font-size:1.9rem}.part p{color:var(--dim)}
blockquote{margin:1.2rem 0 1.6rem;background:var(--soft);border-left:3px solid var(--gold);padding:.8rem 1rem}blockquote p{text-indent:0;text-align:left;margin:0;color:var(--navy);font-size:.95em}.takeaway{border-left:0;border-top:1px solid var(--gold);border-bottom:1px solid var(--gold);background:none;padding:.8rem 0;margin-top:2rem}.takeaway p{font-weight:700}.cite{font-size:.78em;text-decoration:none}.url{overflow-wrap:anywhere}.references p{text-indent:0;text-align:left;font-size:.86em}.tocbox{border:1px solid var(--line);border-radius:10px;padding:.8rem 1rem;margin:2rem 0}.tocbox summary{font-weight:700;color:var(--navy);cursor:pointer}.tocbox ul{list-style:none;padding:0;font-size:.9rem}.tocbox li{margin:.35rem 0}.tocbox li.c{padding-left:1.2em}.tocbox li.p{font-weight:bold;margin-top:.9rem}.tocbox a{text-decoration:none}.foot{text-align:center;color:var(--dim);border-top:1px solid var(--line);padding-top:2rem;margin-top:3rem;font-size:.8rem}.foot p{text-indent:0;text-align:center}.navrow{display:flex;flex-wrap:wrap;justify-content:space-between;gap:1rem;padding:1rem 0;border-top:1px solid var(--line);font-size:.9rem}.note{padding:.7rem 1rem;background:var(--soft);font-size:.87rem}.note p{text-indent:0;margin:0}.formula{text-indent:0;text-align:center}
@media(max-width:600px){body{font-size:17px}.wrap{padding:1rem 1rem 3rem}.hero img{width:170px}.hero h1{font-size:1.5rem}.bar{gap:.55rem;padding:.5rem .7rem}.unit h1,.unit h2{font-size:1.4rem}}
'''
WEB_JS="""(()=>{const key='book25-text-v1';let size=0;try{if(localStorage.getItem(key+'-dark')==='1')document.documentElement.classList.add('dark')}catch(e){};document.getElementById('th').onclick=()=>{document.documentElement.classList.toggle('dark');try{localStorage.setItem(key+'-dark',document.documentElement.classList.contains('dark')?'1':'0')}catch(e){}};document.getElementById('fs').onclick=()=>{size=(size+1)%3;document.body.style.fontSize=[18,20,22][size]+'px'};function progress(){let h=document.documentElement.scrollHeight-innerHeight;document.getElementById('prog').style.width=(h>0?100*scrollY/h:0)+'%'}addEventListener('scroll',progress,{passive:true});progress()})();"""

def toc_html(sections,prefix=''):
    return '<ul>'+''.join('<li class="'+('c' if s['kind']=='chapter' else 'p')+'"><a href="'+prefix+'#'+s['id']+'">'+escaped(s['title'])+'</a></li>' for s in sections)+'</ul>'
def web_shell(title,canonical,inner):
    return '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+escaped(title)+'</title><meta name="description" content="第25号专著数字重编版：六编三十章，五个附录；三律医学哲学与心血管健康教育。"><link rel="canonical" href="https://sdeuniverses.com'+canonical+'"><style>'+WEB_CSS+'</style></head><body><div id="prog"></div><nav class="bar"><a href="'+BASE+'">← 书籍详情</a><a href="'+BASE+'read.html">在线翻页</a><a href="'+BASE+'chapters.html">目录</a><a href="'+BASE+PDFNAME+'">PDF</a><span class="sp"></span><button id="fs" type="button">字号</button><button id="th" type="button">夜间</button></nav><main class="wrap">'+inner+'<footer class="foot"><p>德麦国际专著第25号 · 王德生、秦莉 著 · 数字重编版 v1.0</p><p>医学哲学与健康教育读物，不替代个体诊疗。ISBN '+ISBN+'</p></footer></main><script>'+WEB_JS+'</script></body></html>'

def font(size,bold=False):
    p=Path('/usr/share/fonts/opentype/noto/NotoSerifCJK-'+('Bold' if bold else 'Regular')+'.ttc')
    return ImageFont.truetype(str(p),size,index=2)
def cover_image(path,back=False):
    w,h=1520,2000;im=Image.new('RGB',(w,h),NAVY);d=ImageDraw.Draw(im);cream='#f4eedf';gold='#c6a264'
    def line(txt,y,size=40,fill=cream,bold=False):
        f=font(size,bold);box=d.textbbox((0,0),txt,font=f);x=(w-(box[2]-box[0]))/2;d.text((x,y),txt,font=f,fill=fill)
    d.rectangle((67,67,w-67,h-67),outline='#445675',width=2)
    line('D E M A I   M O N O G R A P H S',118,29,gold)
    line('德麦国际专著 · 第25号',185,36,gold)
    if not back:
        line('心血管健康的',390,116,cream,True);line('发生学重建',558,116,cream,True)
        line('从风险控制到调节能力与生活重建',774,42,gold)
        for j in range(3):
            cx=w/2+(j-1)*135;cy=1160;r=218
            d.ellipse((cx-r,cy-r,cx+r,cy+r),outline=('#76643f' if j!=1 else gold),width=3)
        d.line((390,1160,540,1160,620,1060,710,1260,800,1105,880,1160,1130,1160),fill=gold,width=4)
        line('特征 · 自由 · 幸福',1438,35,gold)
        line('王德生、秦莉 著',1600,47,cream)
        line('The Generative Reconstruction',1725,31,cream)
        line('of Cardiovascular Health',1772,31,cream)
        line('德麦国际有限公司',1880,35,gold)
    else:
        line('保护生命',420,83,cream,True);line('也让生命重新生活',545,83,cream,True)
        for y,t in [(820,'不只问指标是否改善'),(905,'也问行动怎样成为可能'),(990,'以及被保护的生活为何值得继续')]:line(t,y,43,cream)
        d.line((620,1150,900,1150),fill=gold,width=3)
        line('六编三十章 · 五个附录',1240,38,gold)
        line('医学哲学 × 证据边界 × 生活重建',1320,35,gold)
        line('数字重编版 v1.0',1590,38,cream)
        line('ISBN 978-1-970820-33-1',1690,32,gold)
        line('医学哲学与健康教育读物',1810,29,cream)
        line('不替代个体诊疗，不提供自行停药方案',1860,29,cream)
    im.save(path,quality=94,subsampling=0)

def make_docx(path,sections,cover):
    d=Document();sec=d.sections[0];sec.page_width=Mm(190);sec.page_height=Mm(250)
    sec.top_margin=Mm(19);sec.bottom_margin=Mm(18);sec.left_margin=sec.right_margin=Mm(20)
    sec.header_distance=Mm(9);sec.footer_distance=Mm(9)
    for name in ['Normal','Heading 1','Heading 2','Heading 3','Title','Subtitle']:
        st=d.styles[name];st.font.name='Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    normal=d.styles['Normal'];normal.font.size=Pt(11.2);normal.font.color.rgb=RGBColor.from_string('2D2D2D');normal.paragraph_format.line_spacing=Pt(20.6);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.first_line_indent=Pt(22.4)
    for name,size in [('Heading 1',21),('Heading 2',18),('Heading 3',12.2)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('172B49');st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(10);st.paragraph_format.keep_with_next=True;st.paragraph_format.line_spacing=1.45
    p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(cover),width=Mm(145))
    sec.different_first_page_header_footer=True
    hp=sec.header.paragraphs[0];hp.text=TITLE+' · 025';hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:r.font.size=Pt(7);r.font.color.rgb=RGBColor.from_string('888274')
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=fp.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    for s in sections:
        d.add_page_break()
        h=d.add_paragraph(s['title'],'Heading 2' if s['kind']=='chapter' else 'Heading 1')
        if s['kind']=='part':h.paragraph_format.space_before=Pt(70)
        for b in s['blocks']:
            typ=b['type'];p=d.add_paragraph(b['text'],'Heading 3' if typ=='h3' else 'Normal')
            if typ in ('quote','takeaway'):
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Mm(4);p.paragraph_format.right_indent=Mm(3);p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(12)
                pp=p._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F3EEE1');pp.append(sh)
                for r in p.runs:r.font.color.rgb=RGBColor.from_string('172B49');r.bold=typ=='takeaway'
            elif s['kind']=='references':
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(15.5)
                for r in p.runs:r.font.size=Pt(9)
            elif typ=='p':p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    d.core_properties.title=TITLE;d.core_properties.subject=SUB;d.core_properties.author=AUTHORS;d.core_properties.keywords='第25号;三律;SDE;医学哲学;数字重编版v1.0';d.core_properties.comments='由统一章节文本生成的可编辑母稿。医学边界与修订说明见本书附录。'
    d.save(path)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--root',default='.');ap.add_argument('--out',default='build/book25');args=ap.parse_args();root=Path(args.root).resolve();out=(root/args.out).resolve();site=out/'public/books/m/25';site.mkdir(parents=True,exist_ok=True)
    src=root/'tools/book25/manuscript';files=sorted(src.glob('*.md'));assert len(files)==9
    sections,master=parse(files);write(out/'book25-rebuilt-master.md',master);write(site/'book25-rebuilt-master.md',master)
    shutil.copytree(src,out/'manuscript',dirs_exist_ok=True)
    cover_image(site/'cover.jpg');cover_image(site/'backcover.jpg',True)
    stats={'title':TITLE,'authors':[ '王德生','秦莉'],'isbn':ISBN,'edition':'数字重编版 v1.0','version':VERSION,'parts':6,'chapters':30,'appendices':5,'reference_entries':40,'all_han':han(master),'body_han':sum(han(b['text']) for s in sections if s['source'].startswith(('01-','02-','03-','04-','05-','06-')) for b in s['blocks'] if b['type'] in ('p','quote','takeaway')),'paragraphs':sum(len(s['blocks']) for s in sections),'source_files':12,'source_topics':16,'reindex_called':False,'clinical_validation_claimed':False}
    pbody='<div class="cover" id="cover"><img src="cover.jpg" alt="封面"></div>'
    for s in sections:
        if s['kind']=='part' and not any(x['kind']=='part' for x in sections[:sections.index(s)]):pbody+='<section class="toc" id="print-toc"><h1>目录</h1><p style="text-indent:0;color:#8a8578;font-size:9pt">页码为PDF物理页序，含封面与前置页。</p>'+toc_html(sections)+'</section>'
        pbody+=section_html(s)
    pbody+='<div class="cover" id="backcover"><img src="backcover.jpg" alt="封底"></div>'
    print_html='<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>'+TITLE+'</title><meta name="author" content="'+AUTHORS+'"><style>'+PRINT_CSS+'</style></head><body>'+pbody+'</body></html>'
    write(site/'print.html',print_html)
    document=HTML(string=print_html,base_url=str(site)).render();document.write_pdf(str(site/PDFNAME),pdf_tags=True)
    positions={key:i+1 for i,p in enumerate(document.pages) for key in p.anchors}
    assert all(s['id'] in positions for s in sections)
    stats['pdf_pages']=len(document.pages)
    toc=[{'t':'封面','p':'1','g':1,'l':1}]+[{'t':s['title'],'p':str(positions[s['id']]),'g':positions[s['id']],'l':2 if s['kind']=='chapter' else 1} for s in sections]+[{'t':'封底','p':str(len(document.pages)),'g':len(document.pages),'l':1}]
    write(site/'toc.json',json.dumps(toc,ensure_ascii=False,indent=2));write(site/'edition.json',json.dumps(stats,ensure_ascii=False,indent=2))
    make_docx(site/DOCNAME,sections,site/'cover.jpg')
    hero='<header class="hero"><img src="'+BASE+'cover.jpg" alt="心血管健康的发生学重建封面"><h1>'+TITLE+'</h1><p>'+SUB+'</p><p>'+AUTHORS+' 著 · 德麦国际专著第25号 · 数字重编版 v1.0</p></header>'
    caution='<div class="note"><p>本书为医学哲学与健康教育读物。三律框架不作为已验证的诊断系统；不据此自行停药、调整治疗或进行极端训练。教学案例不代表实际疗效。</p></div>'
    full=hero+caution+'<details class="tocbox" open><summary>全书目录 · 六编三十章与五个附录</summary>'+toc_html(sections)+'</details>'+''.join(section_html(s) for s in sections)
    write(site/'text/index.html',web_shell(TITLE+' · 全文阅读',BASE+'text/',full))
    listing=hero+'<h2>章节目录</h2><p>按PDF实际页码建立；全文与翻页共用同一份重编正文。</p><div class="tocbox">'+toc_html(sections,BASE+'text/')+'</div><p><a href="'+BASE+'read.html">进入在线翻页</a> · <a href="'+BASE+'text/">连续阅读全文</a></p>'
    write(site/'chapters.html',web_shell(TITLE+' · 目录',BASE+'chapters.html',listing))
    chapters=[s for s in sections if s['kind']=='chapter']
    for i,s in enumerate(chapters):
        prev=BASE+f'chapters/{i:02d}/' if i else BASE+'text/'
        nex=BASE+f'chapters/{i+2:02d}/' if i<29 else BASE+'text/#'+next(x['id'] for x in sections if x['title'].startswith('结语'))
        nav='<div class="navrow"><a href="'+prev+'">← 上一章 / 全文</a><a href="'+BASE+'chapters.html">总目录</a><a href="'+nex+'">下一章 →</a></div>'
        write(site/f'chapters/{i+1:02d}/index.html',web_shell(s['title']+' · 第25号专著',BASE+f'chapters/{i+1:02d}/',nav+section_html(s,BASE+'text/')+nav))
    template=(root/'public/books/m/220/read.html').read_text(encoding='utf-8')
    reader=template.replace('我的三个宝贝',TITLE).replace('一个理论、一个工具、一声呼喊，与未来的主体经济学',SUB).replace('/books/m/220/','/books/m/25/')
    reader=re.sub(r'(<script[^>]+id="cfg"[^>]*>).*?(</script>)',lambda m:m[1]+json.dumps({'pdf':BASE+PDFNAME,'key':'books-m-25-'+VERSION,'offset':1},ensure_ascii=False)+m[2],reader,flags=re.S)
    reader=re.sub(r'(<script[^>]+id="toc"[^>]*>).*?(</script>)',lambda m:m[1]+json.dumps(toc,ensure_ascii=False)+m[2],reader,flags=re.S)
    reader=reader.replace('fontExtraProperties:true,','fontExtraProperties:true, isEvalSupported:false,')
    reader=re.sub(r'<script src="/(?:taste/wds-companion/wds-read|wds-mode)\.js[^"\n]*" defer></script>','',reader)
    for f in ['pdf.min.js','pdf.worker.min.js']:
        url='https://cdnjs.cloudflare.com/ajax/libs/pdf.js/3.11.174/'+f
        req=urllib.request.Request(url,headers={'User-Agent':'SDE-Book25-Publication/1.0'})
        data=urllib.request.urlopen(req,timeout=60).read();p=site/'assets'/f;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(data)
        reader=reader.replace(url,BASE+'assets/'+f)
    reader=reader.replace('</header>','<a href="'+BASE+'text/" style="color:var(--dim);text-decoration:none;font-size:.82rem">全文</a><a href="'+BASE+PDFNAME+'" style="color:var(--dim);text-decoration:none;font-size:.82rem">PDF</a></header>',1)
    reader=reader.replace('</body>','<noscript><p><a href="'+BASE+'text/">无脚本全文阅读</a> · <a href="'+BASE+PDFNAME+'">直接打开PDF</a></p></noscript></body>')
    write(site/'read.html',reader)
    write(site/'assets/README.txt','PDF.js 3.11.174 distribution from Mozilla/pdf.js, Apache License 2.0. Source: https://github.com/mozilla/pdf.js/tree/v3.11.174 . Fixed trusted book PDF only; isEvalSupported=false. No font files are distributed separately.\n')
    license_url='https://raw.githubusercontent.com/mozilla/pdf.js/v3.11.174/LICENSE'
    try:write(site/'assets/LICENSE-pdfjs.txt',urllib.request.urlopen(license_url,timeout=30).read().decode())
    except Exception:write(site/'assets/LICENSE-pdfjs.txt','Licensed under the Apache License, Version 2.0. https://www.apache.org/licenses/LICENSE-2.0\nCopyright Mozilla Foundation. See source repository for complete notices.\n')
    reference=(root/'public/books/m/220/index.html').read_text(encoding='utf-8');style=re.search(r'<style>(.*?)</style>',reference,re.S).group(1)
    cards=''.join('<li>'+escaped(s['title'])+'</li>' for s in sections if s['kind']=='part')
    summary='从指标、调节能力与生活意义三个层面重建健康理解。全书以三律与SDE组织提问，以医学证据约束临床结论，讨论血脂、血压、血糖以及门诊、家庭、制度和智能工具。'
    landing='<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+TITLE+' · 德麦国际专著第25号</title><meta name="description" content="'+summary+'"><link rel="canonical" href="https://sdeuniverses.com'+BASE+'"><meta property="og:image" content="https://sdeuniverses.com'+BASE+'cover.jpg"><style>'+style+'\n.hero{display:grid;grid-template-columns:250px 1fr;gap:2rem}.hero img{width:100%;height:auto}.actions{display:flex;gap:.6rem;flex-wrap:wrap}.actions a{display:inline-block;border:1px solid var(--gold,#D9A441);padding:.55rem .9rem;border-radius:5px;text-decoration:none}.small{color:var(--dim,#8C949C);font-size:.88rem}.panel{border-top:1px solid var(--line,#232A31);padding-top:1.4rem;margin-top:2rem}.notice{font-size:.9rem;border-left:3px solid var(--gold,#D9A441);padding:.7rem 1rem;background:#151e29}.chapters li{margin:.5rem 0}@media(max-width:720px){.hero{grid-template-columns:1fr}.hero img{max-width:245px;margin:auto;display:block}.wrap{padding-top:1.2rem}}</style></head><body><main class="wrap"><nav><a href="/books/">← 专著书架</a> · <a href="/monographs/">专著栏目</a></nav><p class="small">DEMAI MONOGRAPHS / 025 / 数字重编版 v1.0</p><div class="hero"><img src="cover.jpg" alt="本书封面"><div><h1>'+TITLE+'</h1><p style="color:var(--gold)">'+SUB+'</p><p>'+AUTHORS+' 著</p><p class="small">ISBN '+ISBN+'<br>六编三十章 · 五个附录 · 40项参考资料<br>PDF '+str(stats['pdf_pages'])+' 页 · 正文 '+format(stats['body_han'],',')+' 汉字</p><div class="actions"><a href="read.html">在线翻页</a><a href="text/">全文阅读</a><a href="chapters.html">章节目录</a><a href="'+PDFNAME+'">PDF下载</a><a href="'+DOCNAME+'">Word母稿</a></div></div></div><section class="panel"><h2>保护生命，也让生命重新生活</h2><p>'+summary+'</p><p>本版由11篇独立论文与整书初稿所含16项材料重编，不是原稿逐句转录。重组长段落，接通论证，增加案例与反例，并逐项区分理论、证据与实践建议。</p><div class="notice">三律框架不作为已验证的诊断或治疗系统。书中教学案例为构拟，不代表实际疗效；不据此自行停药、改变治疗或进行极端训练。实质修订见附录四。</div></section><section class="panel"><h2>全书结构</h2><ul class="chapters">'+cards+'</ul><p><a href="chapters.html">展开完整目录 →</a></p></section><section class="panel"><h2>出版与版本</h2><p>德麦国际有限公司 · 第25号专著<br>沿用既有书名、作者顺序与ISBN；第220号仅作版式和阅读体验参照。</p><p class="small">本版未另行编造定价，亦不表示已取得外部临床审稿或新版纸书登记。正文与证据校订截至2026年10月2日。</p><p><a href="book25-rebuilt-master.md">可编辑Markdown全文</a> · <a href="edition.json">版本统计</a></p></section><section class="panel"><h2>原有专题入口</h2><p class="small">以下为保留的历史理论专题，不作为本次重编的临床证据。</p><p><a href="articles/01/">三律生理学入门</a> · <a href="articles/02/">器官三律意义工程学入门</a> · <a href="articles/03/">三律组织学入门</a></p></section><footer class="panel small">© '+AUTHORS+' · 德麦国际专著第25号</footer></main></body></html>'
    write(site/'index.html',landing)
    # Verify every text unit and all local reader targets before producing a publication manifest.
    doc=fitz.open(site/PDFNAME);pdftext=''.join(p.get_text() for p in doc)
    assert len(doc)==stats['pdf_pages'] and len(doc)>60
    normalized=re.sub(r'\s+','',pdftext)
    for s in sections:assert re.sub(r'\s+','',s['title']) in normalized,s['title']
    assert '我的三个宝贝' not in reader and 'books/m/220/' not in reader
    assert 'isEvalSupported:false' in reader
    assert sum(1 for p in site.rglob('index.html') if '/chapters/' in str(p))==30
    manifest={'edition':VERSION,'stats':stats,'reindex_called':False,'files':[{ 'path':p.relative_to(out).as_posix(),'bytes':p.stat().st_size,'sha256':digest(p)} for p in sorted((out/'public').rglob('*')) if p.is_file()]}
    write(out/'publication-manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2));write(out/'toc-physical-pages.json',json.dumps(toc,ensure_ascii=False,indent=2))
    print(json.dumps(stats,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
