"""Finalize author-confirmed ISBN for monograph 289; retain the accepted v1.1 text and layout."""
from pathlib import Path
import argparse, hashlib, io, json, re, shutil, subprocess, zipfile
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from bs4 import BeautifulSoup
from reportlab.graphics.barcode import createBarcodeDrawing
from reportlab.graphics import renderPDF
from reportlab.lib.units import mm
import fitz
ISBN='979-8-90690-376-1'
REV='20261002-isbn376'
BASE_SHA='749eee90c4e4a74fd0523bc11e00626f67bd5dbbd973b66009f33525a4069bde'
BOOK='education-subject-rebirth'
def sha(data): return hashlib.sha256(data).hexdigest()
def sig_doc(d):
    start=next(i for i,p in enumerate(d.paragraphs) if p.text.startswith('导读') and p.style.name=='Heading 1')
    return [p._p.xml for p in d.paragraphs[start:] if p.style.name!='PubBlank']
def sig_html(data):
    s=BeautifulSoup(data,'html.parser')
    return [(p.get('data-source-paragraph'),p.get_text()) for p in s.select('[data-source-paragraph]')]
def prepare(src,out,work):
    m=json.loads((src/'publication-manifest.json').read_text())
    assert m['release_sha256']==BASE_SHA, 'Source edition changed; reconcile instead of overwriting.'
    assert m['number']==289 and m['price']==20
    digits=ISBN.replace('-',''); assert len(digits)==13
    assert (-sum(int(x)*(1 if i%2==0 else 3) for i,x in enumerate(digits[:-1])))%10==int(digits[-1])
    for n,e in m['files'].items(): assert sha((src/n).read_bytes())==e['sha256'], n
    shutil.copytree(src,out,dirs_exist_ok=True); work.mkdir(parents=True,exist_ok=True)
    assets=work/'assets';assets.mkdir(exist_ok=True)
    font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    with zipfile.ZipFile(src/'downloads/education-subject-rebirth.docx') as z:
        for name,imname in [('cover','image1.png'),('backcover','image2.png')]:
            im=Image.open(io.BytesIO(z.read('word/media/'+imname))).convert('RGB')
            sx,sy=im.width/170,im.height/240; draw=ImageDraw.Draw(im)
            if name=='cover':
                draw.text((157*sx,219*sy),'ISBN '+ISBN,font=ImageFont.truetype(font_path,round(2.6*sy)),fill='#F2EDDF',anchor='rt')
            else:
                left,top,width,height=104,198,53,23
                draw.rectangle((left*sx,top*sy,(left+width)*sx,(top+height)*sy),fill='white')
                draw.text(((left+width/2)*sx,(top+1)*sy),'ISBN '+ISBN,font=ImageFont.truetype(font_path,round(2.4*sy)),fill='#101823',anchor='mt')
                barcode=createBarcodeDrawing('EAN13',value=digits[:12],humanReadable=True,barHeight=16*mm,barWidth=.33*mm,fontSize=8)
                bp=fitz.open(stream=renderPDF.drawToString(barcode),filetype='pdf')
                pix=bp[0].get_pixmap(matrix=fitz.Matrix(sx*25.4/72,sy*25.4/72),alpha=False)
                bi=Image.frombytes('RGB',(pix.width,pix.height),pix.samples);bp.close()
                bi.thumbnail((round(49*sx),round(18*sy)),Image.Resampling.LANCZOS)
                im.paste(bi,(round((left+width/2)*sx-bi.width/2),round((top+4)*sy)))
            im.save(assets/(name+'.png'),optimize=True)
            im.save(out/(name+'.jpg'),quality=95,subsampling=0,optimize=True)
    dp=src/'downloads/education-subject-rebirth.docx';d=Document(dp);prior=sig_doc(d)
    repl={'ISBN\u3000待确认':'ISBN\u3000'+ISBN,
          '第289卷与定价由作者确认，ISBN另行核对。印刷校样供版式复核；纸张、出血、书脊与装订参数须由承印厂另行确认。':
          '第289卷、ISBN及定价由作者确认。印刷校样供版式复核；纸张、出血、书脊与装订参数须由承印厂另行确认。'}
    changes=[]
    for p in d.paragraphs:
        if p.text in repl:
            old=p.text;new=repl[old]
            if p.runs:
                p.runs[0].text=new
                for r in p.runs[1:]:r.text=''
            else:p.add_run(new)
            changes.append({'before':old,'after':new})
    assert len(changes)==2, changes
    for rel in d.part.rels.values():
        if rel.reltype.endswith('/image'):
            name=Path(rel.target_ref).name
            if name=='image1.png':rel.target_part._blob=(assets/'cover.png').read_bytes()
            if name=='image2.png':rel.target_part._blob=(assets/'backcover.png').read_bytes()
    d.core_properties.subject='德麦国际专著第289卷 | ISBN '+ISBN+' | US$20.00'
    d.core_properties.keywords='SDE;教育发生学;复合主体;ISBN '+ISBN
    d.core_properties.comments='校订v1.1，ISBN确认版。书号由作者确认；正文与附录未改。'
    dest=out/'downloads/education-subject-rebirth.docx';d.save(dest)
    assert sig_doc(Document(dest))==prior, 'Manuscript XML changed'
    (work/'word-change-audit.json').write_text(json.dumps({'isbn':ISBN,'revision':REV,'text_changes':changes,'body_and_appendix_xml_unchanged':True},ensure_ascii=False,indent=2))
    return dest

def finish(src,out,work,rendered):
    newpage=fitz.open(rendered);idx=2 if len(newpage)>1 else 0
    assert ISBN in newpage[idx].get_text() and 'US$20.00' in newpage[idx].get_text()
    pdf=fitz.open(src/'downloads/education-subject-rebirth-print.pdf');assert len(pdf)==333
    bookmarks=pdf.get_toc();links=sum(len(p.get_links()) for p in pdf)
    for i in [0,2,332]:
        p=pdf[i];p.add_redact_annot(p.rect,fill=False);p.apply_redactions(images=2,graphics=2)
        if i==2:p.show_pdf_page(p.rect,newpage,idx,keep_proportion=False)
        else:p.insert_image(p.rect,filename=str(work/'assets'/('cover.png' if i==0 else 'backcover.png')),keep_proportion=False)
    newpage.close()
    meta=pdf.metadata;meta['subject']='德麦国际专著第289卷 | ISBN '+ISBN+' | US$20.00 | 校订v1.1 ISBN确认版';meta['keywords']='SDE;教育发生学;ISBN '+ISBN
    pdf.set_metadata(meta)
    printed=out/'downloads/education-subject-rebirth-print.pdf';printed.unlink()
    pdf.save(printed,garbage=4,deflate=True,no_new_id=True);pdf.close()
    read=fitz.open(printed)
    for p in list(read)[1:-1]:p.draw_rect(p.rect,color=None,fill=(.9843,.9725,.9412),overlay=False)
    reader=out/'downloads/education-subject-rebirth-reader.pdf';reader.unlink();read.save(reader,garbage=4,deflate=True,no_new_id=True);read.close()
    pdf_checks=[]
    for name in ['education-subject-rebirth-print.pdf','education-subject-rebirth-reader.pdf']:
        old=fitz.open(src/'downloads'/name);new=fitz.open(out/'downloads'/name)
        assert len(new)==333 and new.get_toc()==bookmarks
        assert sum(len(p.get_links()) for p in new)==links
        textdiff=[];pixdiff=[]
        for i in range(333):
            if old[i].get_text()!=new[i].get_text():textdiff.append(i+1)
            if old[i].get_pixmap(matrix=fitz.Matrix(.8,.8),alpha=False).samples!=new[i].get_pixmap(matrix=fitz.Matrix(.8,.8),alpha=False).samples:pixdiff.append(i+1)
        assert set(textdiff)<={1,3,333},textdiff
        assert set(pixdiff)<={1,3,333},pixdiff
        assert ISBN in new[2].get_text() and '待确认' not in new[2].get_text()
        pdf_checks.append({'file':name,'pages':333,'bookmarks':len(bookmarks),'internal_links':links,'changed_text_pages':textdiff,'changed_pixel_pages':pixdiff})
        for i in [0,2,163,332]:new[i].get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(work/(name+f'-page{i+1}.png'))
        old.close();new.close()
    for p in out.rglob('*.html'):
        raw=p.read_text();before=sig_html(raw);s=BeautifulSoup(raw,'html.parser')
        for t in list(s.find_all(string=lambda x:x and ('待确认' in x or '待作者确认' in x))):
            text=str(t).replace('ISBN\u3000待确认','ISBN\u3000'+ISBN).replace('ISBN 待确认','ISBN '+ISBN).replace('ISBN待确认','ISBN '+ISBN)
            if text.strip()=='待确认' and t.parent.name=='dd':
                dt=t.parent.find_previous_sibling('dt')
                if dt and dt.get_text().strip()=='ISBN':text=ISBN
            if text!=str(t):t.replace_with(text)
        tag=s.find('meta',attrs={'name':'isbn'})
        if not tag:tag=s.new_tag('meta',attrs={'name':'isbn'});s.head.append(tag)
        tag['content']=ISBN
        for t in s.select('script[type="application/ld+json"]'):
            j=json.loads(t.string);j['isbn']=ISBN;t.string=json.dumps(j,ensure_ascii=False)
        for im in s.select('img[src]'):
            if im['src'].split('?')[0].endswith(('cover.jpg','backcover.jpg')):im['src']=im['src'].split('?')[0]+'?v='+REV
        for a in s.select('a[href]'):
            if 'education-subject-rebirth' in a['href'] and a['href'].split('?')[0].endswith(('.docx','.pdf')):a['href']=a['href'].split('?')[0]+'?v='+REV
        result=str(s)
        if p.name=='read.html':
            result=result.replace('downloads/education-subject-rebirth-reader.pdf"','downloads/education-subject-rebirth-reader.pdf?v='+REV+'"').replace("downloads/education-subject-rebirth-reader.pdf'","downloads/education-subject-rebirth-reader.pdf?v="+REV+"'")
        assert before==sig_html(result),p
        assert 'ISBN待确认' not in result and 'ISBN 待确认' not in result and 'ISBN\u3000待确认' not in result,p
        assert '979-8-90690-373-0' not in result and '979-8-90690-365-5' not in result
        p.write_text(result)
    m=json.loads((src/'publication-manifest.json').read_text());m.update(isbn=ISBN,isbn_status='author_confirmed',metadata_revision=REV)
    m['files']={str(p.relative_to(out)):{'sha256':sha(p.read_bytes()),'bytes':p.stat().st_size} for p in sorted(out.rglob('*')) if p.is_file() and p.name not in ['publication-manifest.json','release-check.json']}
    m['release_sha256']=sha(json.dumps(m['files'],sort_keys=True).encode())
    (out/'publication-manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2))
    audit={'volume':289,'isbn':ISBN,'isbn_checksum_valid':True,'isbn_basis':'author_confirmation','price_usd':20,'edition':'校订v1.1 ISBN确认版','revision':REV,'files':len(m['files']),'manuscript_xml_unchanged':True,'html_source_blocks_unchanged':True,'pdf_checks':pdf_checks,'release_sha256':m['release_sha256'],'reindex_triggered':False}
    (out/'release-check.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2));(work/'build-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2));print(json.dumps(audit,ensure_ascii=False,indent=2))
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('source',type=Path);ap.add_argument('output',type=Path);ap.add_argument('work',type=Path);ap.add_argument('--rendered',type=Path)
    a=ap.parse_args();a.source=a.source.resolve();a.output=a.output.resolve();a.work=a.work.resolve()
    if a.rendered:finish(a.source,a.output,a.work,a.rendered)
    else:prepare(a.source,a.output,a.work)
