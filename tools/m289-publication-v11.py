"""Publish monograph 289 v1.1 from the authorized edition. Preserve all manuscript text."""
from pathlib import Path
import json,re,shutil,hashlib,sys
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Pt
from bs4 import BeautifulSoup
from reportlab.graphics.barcode import createBarcodeDrawing
from reportlab.graphics import renderPDF
import fitz
TITLE='AI时代教育对象的重生';ISBN='979-8-90690-365-5';NO=289;VERSION='1.1'
OLD='/books/education-subject-rebirth/';NEW='/books/m/289/'
def digest(data):return hashlib.sha256(data).hexdigest()
def cover_assets(source,out):
 out.mkdir(parents=True,exist_ok=True)
 reg='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc';bold='/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
 for stem in ['cover','backcover']:
  im=Image.open(source/(stem+'.jpg')).convert('RGB');W,H=im.size;sx,sy=W/170,H/240;draw=ImageDraw.Draw(im);gold='#D6B763';cream='#F2EDDF'
  def font(mm,b=False):return ImageFont.truetype(bold if b else reg,round(mm*sy),index=2)
  def txt(x,y,t,mm=2.5,color=gold,b=False,anchor=None):draw.text((round(x*sx),round(y*sy)),t,font=font(mm,b),fill=color,anchor=anchor)
  if stem=='cover':
   txt(157,14,'289',10,gold,True,anchor='rt');txt(157,26,'德麦国际专著',2.5,gold,anchor='rt')
   draw.line([(round(116*sx),round(34*sy)),(round(157*sx),round(34*sy))],fill=gold,width=2)
   txt(157,226,'ISBN '+ISBN,2.45,cream,anchor='rt');txt(157,230,'第289卷 · 2026年10月',2.15,gold,anchor='rt')
  else:
   txt(157,183,'DEMAI  /  289',3.2,gold,True,anchor='rt')
   digits=ISBN.replace('-','');assert (-sum(int(c)*(1 if i%2==0 else 3) for i,c in enumerate(digits[:-1])))%10==int(digits[-1])
   drawing=createBarcodeDrawing('EAN13',value=digits[:12],humanReadable=True,barHeight=39,barWidth=1.05)
   p=fitz.open(stream=renderPDF.drawToString(drawing),filetype='pdf');pix=p[0].get_pixmap(matrix=fitz.Matrix(5,5),alpha=False)
   bi=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);p.close()
   box=(round(104*sx),round(202*sy),round(157*sx),round(225*sy));draw.rectangle(box,fill='#FFFFFF')
   draw.text((box[0]+round(2*sx),box[1]+round(1*sy)),'ISBN '+ISBN,font=font(2.45),fill='#101823')
   im.paste(bi.resize((round(48*sx),round(15*sy)),Image.Resampling.LANCZOS),(box[0]+round(2.5*sx),box[1]+round(6*sy)))
   draw=ImageDraw.Draw(im);txt(157,228,'定价 US$20.00',2.7,cream,anchor='rt')
  im.save(out/(stem+'.png'),optimize=True);im.save(out/(stem+'.jpg'),quality=95,subsampling=0,optimize=True)
def edit_docx(src,dest,assets):
 d=Document(src)
 replacements={
 '德麦国际专著  /  教育与学习':'德麦国际专著第289卷  /  教育与学习',
 '数字阅读版  ·  2026年10月':'第1版 · 校订 v1.1  /  2026年10月',
 '版本　2026年10月数字阅读版 · v1.0':'版本　2026年10月第1版 · 校订 v1.1',
 '阅读入口　SDEUniverses.com · 专著书架':'阅读入口　SDEUniverses.com/books/m/289/',
 '本书的正式专著卷号、ISBN与纸书定价尚未确认，本版不作编造或借用。印刷校样供版式复核；书脊、出血及具体印厂参数应在制作纸书时另行确认。':'第289卷、ISBN及定价由作者确认。印刷校样供版式复核；纸张、出血、书脊与装订参数须由承印厂另行确认。',
 '本书的理论问题与作者的SDE研究及相关平台实践存在直接联系。这里的作者介绍不构成对理论效果的独立认证；书中构造性情境与待执行研究均保持其原有材料身份。':'本书沿着作者的理论研究与知识生态实践展开，以教育为切入点，讨论知识的生成、传播与更新，以及人如何参与人—智能器具关系的形成与重构。'}
 changed=[]
 for i,p in enumerate(d.paragraphs):
  if p.text in replacements:
   old=p.text;p.text=replacements[old];changed.append({'paragraph':i,'before':old,'after':p.text})
 target=next(p for p in d.paragraphs if p.text=='开本　170 mm × 240 mm')
 for t in ['卷号　德麦国际专著第289卷','ISBN　979-8-90690-365-5','定价　US$20.00']:
  p=target.insert_paragraph_before(t,style='PubMeta');p.paragraph_format.space_after=Pt(6)
 for rel in d.part.rels.values():
  if rel.reltype.endswith('/image'):
   name=Path(rel.target_ref).name
   if name=='image1.png':rel.target_part._blob=(assets/'cover.png').read_bytes()
   if name=='image2.png':rel.target_part._blob=(assets/'backcover.png').read_bytes()
 cp=d.core_properties;cp.title=TITLE;cp.author='王德生';cp.subject='第289卷｜'+ISBN+'｜US$20.00';cp.version=VERSION
 cp.keywords='教育发生学;SDE;复合主体;德麦国际专著第289卷;ISBN '+ISBN
 cp.comments='校订v1.1。出版标识、封面封底及阅读呈现更新；不作为独立学术外审。'
 dest.parent.mkdir(parents=True,exist_ok=True);d.save(dest)
 def signature(doc):
  start=next(i for i,p in enumerate(doc.paragraphs) if p.style.name=='Heading 1' and p.text.startswith('前言'))
  return [p._p.xpath('.//w:t/text()')+p._p.xpath('.//m:t/text()') for p in doc.paragraphs[start:] if p.style.name!='PubBlank']
 assert signature(Document(src))==signature(Document(dest)),'Manuscript changed'
 return changed
READER_CSS_ADD='''
/* 289 · 校订版：阅读层次与可访问性，不改正文 */
html{scroll-behavior:smooth;scroll-padding-top:110px}body{line-break:strict}p{overflow-wrap:break-word;text-justify:inter-ideograph}.ref-item p,.reference p{overflow-wrap:anywhere}
.pub-ident{display:flex;flex-wrap:wrap;gap:.6rem 1.4rem;justify-content:center;padding:1rem;border-block:1px solid var(--line);font:13px/1.7 'Noto Sans CJK SC',sans-serif;color:var(--dim)}
.publication-grid{display:grid;grid-template-columns:5rem 1fr;gap:.5rem 1rem;font-size:.95em}.publication-grid dt{color:var(--gold)}.publication-grid dd{margin:0;text-align:left}
.toc-search{display:block;width:100%;padding:.65rem .8rem;margin:.8rem 0;border:1px solid var(--line);border-radius:6px;background:var(--paper);color:var(--ink);font:15px/1.6 sans-serif}.toc-empty{font-size:.9em;text-indent:0;color:var(--dim)}
#reading-status{color:var(--dim);font-size:12px;max-width:16rem;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.resume-card{position:fixed;right:1rem;bottom:1rem;z-index:20;max-width:min(370px,calc(100vw - 2rem));padding:.8rem 1rem;background:var(--paper);border:1px solid var(--gold);border-radius:9px;box-shadow:0 8px 35px #0002;font:14px/1.6 sans-serif}.resume-card button{margin:.3rem .5rem .1rem 0;padding:.4rem .6rem;border:1px solid var(--line);border-radius:5px;color:var(--navy);background:var(--soft)}
.chapter>h2,.chapter h3{text-wrap:pretty}.chapter h3{scroll-margin-top:6rem}.toc [hidden]{display:none!important}.skip-reader{position:absolute;left:1rem;top:-10rem;z-index:40;padding:.8rem;background:var(--paper)}.skip-reader:focus{top:1rem}
@media(max-width:600px){.bar button{min-height:40px;min-width:40px}.bar{align-items:center}.pub-ident{font-size:12px;padding:.8rem .3rem;gap:.4rem .7rem}#reading-status{max-width:100%;flex-basis:100%}.chapter{margin-top:3rem}.publication-grid{grid-template-columns:4rem 1fr;gap:.5rem}.toc-search{font-size:16px}}
@media(prefers-reduced-motion:reduce){html{scroll-behavior:auto}}@media print{.resume-card,.skip-reader,.toc-search,#reading-status{display:none!important}}
'''
READER_JS='''(function(){'use strict';
const root=document.documentElement,key='education-subject-rebirth',path=location.pathname.replace('/books/m/289/','/books/education-subject-rebirth/'),store=key+path;
const get=k=>{try{return localStorage.getItem(k)}catch(e){return null}},set=(k,v)=>{try{localStorage.setItem(k,v)}catch(e){}};
if(get(key+'-dark')==='1')root.classList.add('dark');let n=Number(get(key+'-font'));if(n>=15&&n<=24)root.style.setProperty('--text-size',n+'px');
const theme=document.getElementById('theme');const syncTheme=()=>theme?.setAttribute('aria-pressed',root.classList.contains('dark')?'true':'false');syncTheme();theme?.addEventListener('click',()=>{root.classList.toggle('dark');set(key+'-dark',root.classList.contains('dark')?'1':'0');syncTheme()});
['larger','smaller'].forEach(id=>document.getElementById(id)?.addEventListener('click',()=>{n=parseFloat(getComputedStyle(root).getPropertyValue('--text-size'))||18;n=Math.max(15,Math.min(24,n+(id==='larger'?1:-1)));root.style.setProperty('--text-size',n+'px');set(key+'-font',n)}));
const sections=[...document.querySelectorAll('main section[id]')];let timer,active=null;
function progress(){const max=root.scrollHeight-innerHeight,bar=document.getElementById('progress');if(bar)bar.style.width=(max>0?100*scrollY/max:0)+'%';for(const s of sections){if(s.getBoundingClientRect().top<=140)active=s;else break}const status=document.getElementById('reading-status');if(status)status.textContent=(active?.querySelector('h2')?.textContent||'阅读')+' · '+Math.round(max>0?100*scrollY/max:0)+'%';clearTimeout(timer);timer=setTimeout(()=>{set(store,String(scrollY));if(active)set(store+'-anchor',JSON.stringify({id:active.id,offset:Math.max(0,scrollY-active.offsetTop)}))},600)}
window.addEventListener('scroll',progress,{passive:true});window.addEventListener('resize',progress,{passive:true});
if(!location.hash){let saved=null;try{saved=JSON.parse(get(store+'-anchor')||'null')}catch(e){}const y=Number(get(store));if(y>500){const card=document.createElement('aside');card.className='resume-card';card.setAttribute('aria-label','阅读进度');card.textContent='已保存上次阅读位置。';const resume=document.createElement('button');resume.textContent='继续上次阅读';resume.onclick=()=>{const target=saved?.id&&document.getElementById(saved.id);scrollTo(0,target?target.offsetTop+(saved.offset||0):y);card.remove()};const close=document.createElement('button');close.textContent='从当前位置读';close.onclick=()=>card.remove();card.append(resume,close);document.body.append(card)}}
const search=document.querySelector('.toc-search');if(search){const lis=[...search.closest('.toc').querySelectorAll('li')],empty=search.closest('.toc').querySelector('.toc-empty');search.addEventListener('input',()=>{const q=search.value.trim().toLocaleLowerCase();let count=0;for(const li of lis){li.hidden=!!q&&!li.textContent.toLocaleLowerCase().includes(q);if(!li.hidden)count++}if(empty)empty.hidden=count>0})}
const max=root.scrollHeight-innerHeight;const bar=document.getElementById('progress');if(bar)bar.style.width=(max>0?100*scrollY/max:0)+'%';
})();'''
PUB_HTML=f'''<section class="front-sec" id="publication"><div class="chap-label">DEMAI INTERNATIONAL PRESS / 289</div><h2>出版信息</h2><dl class="publication-grid"><dt>书名</dt><dd>{TITLE}</dd><dt>副题</dt><dd>从封闭学生到人—智能器具复合主体</dd><dt>著者</dt><dd>王德生</dd><dt>出版</dt><dd>德麦国际出版社</dd><dt>卷号</dt><dd>德麦国际专著第289卷</dd><dt>ISBN</dt><dd>{ISBN}</dd><dt>定价</dt><dd>US$20.00</dd><dt>版本</dt><dd>2026年10月第1版 · 校订 v1.1</dd><dt>规格</dt><dd>170 × 240毫米 · 六编三十章 · 333个PDF页面（含封面封底）</dd></dl><p class="small">全文在线开放阅读；定价不改变本页的开放阅读安排。纸张、出血、书脊及装订参数由承印厂另行确认。</p></section>'''
def signature_html(raw):
 s=BeautifulSoup(raw,'html.parser');return [(p.get('data-source-paragraph'),p.get_text()) for p in s.select('[data-source-paragraph]')]
def site_pages(source,out):
 shutil.copytree(source,out,dirs_exist_ok=True)
 (out/'reader.css').write_text((source/'reader.css').read_text()+READER_CSS_ADD);(out/'reader.js').write_text(READER_JS)
 author_old='本书的理论问题与作者的SDE研究及相关平台实践存在直接联系。这里的作者介绍不构成对理论效果的独立认证；书中构造性情境与待执行研究均保持其原有材料身份。'
 author_new='本书沿着作者的理论研究与知识生态实践展开，以教育为切入点，讨论知识的生成、传播与更新，以及人如何参与人—智能器具关系的形成与重构。'
 for f in list(out.rglob('*.html')):
  raw=f.read_text();before=signature_html(raw)
  text=raw.replace(OLD,NEW).replace('?v=1.0','?v=1.1').replace('v1.0','v1.1').replace(author_old,author_new)
  s=BeautifulSoup(text,'html.parser')
  if s.title and '第289卷' not in s.title.get_text():s.title.string=s.title.get_text()+' · 第289卷'
  for k,v in [('book_no','289'),('isbn',ISBN),('edition','v1.1')]:
   tag=s.find('meta',attrs={'name':k})
   if not tag:tag=s.new_tag('meta',attrs={'name':k});s.head.append(tag)
   tag['content']=v
  for t in s.select('script[type="application/ld+json"]'):
   j=json.loads(t.string or '{}');j.update(isbn=ISBN,bookEdition='2026年10月第1版 · 校订v1.1',numberOfPages=333)
   j['identifier']={'@type':'PropertyValue','propertyID':'德麦国际专著卷号','value':'289'}
   j['offers']={'@type':'Offer','price':'20.00','priceCurrency':'USD','url':'https://sdeuniverses.com'+NEW};t.string=json.dumps(j,ensure_ascii=False)
  hero=s.select_one('main .hero')
  if hero and f.name!='read.html' and f.relative_to(out).as_posix()!='index.html':
   hero.append(BeautifulSoup(f'<div class="pub-ident"><span>德麦国际专著第289卷</span><span>ISBN {ISBN}</span><span>US$20.00</span><span>校订 v1.1</span></div>','html.parser').div)
  toc=s.select_one('details.toc')
  if toc:
   inp=s.new_tag('input',attrs={'class':'toc-search','type':'search','placeholder':'查找章节或主题…','aria-label':'查找目录','autocomplete':'off'});toc.summary.insert_after(inp)
   empty=s.new_tag('p',attrs={'class':'toc-empty','hidden':''});empty.string='未找到对应章节，请尝试较短的关键词。';toc.append(empty)
  if f.relative_to(out).as_posix()=='text/index.html':
   s.find(id='author').insert_before(BeautifulSoup(PUB_HTML,'html.parser').section)
   ol=toc.find(['ol','ul']) if toc else None
   if ol:
    li=s.new_tag('li');a=s.new_tag('a',href='#publication');a.string='出版信息 · 第289卷';li.append(a);ol.insert(0,li)
  bar=s.select_one('.bar')
  if bar:
   bar['aria-label']='阅读工具栏';bar['role']='navigation';status=s.new_tag('span',id='reading-status');status['aria-live']='off';status.string='第289卷 · 校订 v1.1';bar.append(status)
   for b in bar.select('button'):
    b['type']='button'
    if b.get('id')=='theme':b['aria-label']='切换夜间模式'
    elif b.get('id')=='larger':b['aria-label']='增大字号'
    elif b.get('id')=='smaller':b['aria-label']='减小字号'
  main=s.find('main')
  if main and f.name!='read.html':
   main['id']='main-content';skip=s.new_tag('a',href='#main-content',attrs={'class':'skip-reader'});skip.string='跳到正文';s.body.insert(0,skip)
  for a in s.select('a[target="_blank"]'):a['rel']='noopener noreferrer'
  if f.relative_to(out).as_posix()=='index.html':
   s.select_one('.crumb').append(' · 第289卷');meta=s.select_one('.meta')
   for textline in ['卷号　德麦国际专著第289卷','ISBN　'+ISBN,'定价　US$20.00']:
    p=s.new_tag('p');p.string=textline;meta.append(p)
   for p in meta.find_all('p'):
    if p.get_text().startswith('版本'):p.string='版本　2026年10月第1版 · 校订 v1.1'
   note=s.select_one('.notes')
   if note:note.string='书中理论建构、构造性课堂情境与待执行研究均保留其材料边界。出版编校不替代独立学术外审。本版采用作者确认的第289卷、ISBN 979-8-90690-365-5及US$20.00定价；全文继续开放阅读。'
   for i,c in enumerate(s.select('.card')):
    h=c.find('h3');a=s.new_tag('a',href=NEW+'text/#part'+str(i+1));a.string=h.get_text();h.clear();h.append(a)
   s.style.string=(s.style.string or '')+'''\n.meta{display:grid;grid-template-columns:1fr 1fr;gap:.2rem 1rem}.meta p:nth-child(2),.meta p:nth-child(4),.meta p:nth-child(5){grid-column:1/-1}.card{transition:transform .18s,border-color .18s}.card:hover{transform:translateY(-2px);border-color:#a88945}.card a{color:inherit}.buttons a{min-height:44px;display:inline-flex;align-items:center}.skip-reader{position:absolute;top:-200px;left:1rem;background:#11171f;padding:.6rem;z-index:100}.skip-reader:focus{top:1rem}a:focus-visible{outline:3px solid #d9a441;outline-offset:4px}.hero{padding-bottom:1.6rem;border-bottom:1px solid #293039}.sub{line-height:1.7}.notes{line-height:1.9}@media(max-width:760px){.meta{grid-template-columns:1fr}.meta p{grid-column:auto!important}.hero img{max-width:240px}.wrap{padding-top:1.3rem}}@media(prefers-reduced-motion:reduce){.card{transition:none}}'''
  result=str(s);assert before==signature_html(result),('Manuscript modified',f);f.write_text(result)
 p=out/'read.html';s=p.read_text()
 s=s.replace('<span class="ttl">'+TITLE+'</span>','<span class="ttl">第289卷 · '+TITLE+'</span>')
 s=s.replace('fontExtraProperties:true, isOffscreenCanvasSupported:false','fontExtraProperties:true, isEvalSupported:false, isOffscreenCanvasSupported:false')
 s=s.replace("$('pt').textContent = Math.max(1,total-CFG.offset+1);","$('pt').textContent = Math.max(1,total-CFG.offset); $('pi').max=String(total-CFG.offset); $('pageState').textContent=cur===1?'封面 · PDF 1/'+total:(cur===total?'封底 · PDF '+total+'/'+total:'PDF '+cur+'/'+total);")
 s=s.replace('var printed = Math.max(1, cur-CFG.offset+1);','var printed = Math.max(1,Math.min(total-CFG.offset, cur-CFG.offset+1));').replace('>矢量</button>','>高清</button>')
 soup=BeautifulSoup(s,'html.parser');pg=soup.select_one('span.pg');pg.insert(0,'书内 ')
 state=soup.new_tag('span',id='pageState');state['aria-live']='polite';state.string='PDF · 333页（含封面封底）';soup.select_one('footer.bot').append(state)
 soup.find(id='pi')['aria-label']='跳到书内印刷页码';soup.find(id='pi')['max']='331';soup.find(id='q')['aria-label']='搜索全书正文'
 for b in soup.select('button'):
  b['type']='button'
  if b.get('title'):b['aria-label']=b['title']
 soup.style.string=(soup.style.string or '')+'''\n#pageState{font:11px/1.5 sans-serif;color:var(--dim);white-space:nowrap}.stage{overflow:auto;align-items:safe center;justify-items:safe center}button:focus-visible,a:focus-visible,input:focus-visible{outline:2px solid var(--gold);outline-offset:3px}@media(max-width:700px){#pageState{position:absolute;bottom:60px;right:10px;padding:3px 7px;background:var(--bg);border-radius:5px}.bot button{min-width:38px;min-height:38px}.top .ttl{max-width:105px}.drawer{bottom:65px}}'''
 p.write_text(str(soup))
def finalize_pdf(rendered,source_pdf,dest,assets):
 src=fitz.open(source_pdf);front=fitz.open(rendered);pdf=fitz.open(source_pdf)
 assert len(pdf)==333 and len(front)>=4,'Invalid source/front pagination'
 for i in [1,2,3]:
  page=pdf[i];page.add_redact_annot(page.rect,fill=(1,1,1));page.apply_redactions(images=2,graphics=2);page.show_pdf_page(page.rect,front,i,keep_proportion=False)
 for i,name in [(0,'cover.png'),(332,'backcover.png')]:
  page=pdf[i];page.add_redact_annot(page.rect,fill=(1,1,1));page.apply_redactions(images=2,graphics=2);page.insert_image(page.rect,filename=str(assets/name),keep_proportion=False)
 front.close();assert len(pdf.get_toc())==240
 norm=lambda p:re.sub(r'\s+','',p.get_text())
 differences=[i+1 for i in range(333) if norm(src[i])!=norm(pdf[i])];assert set(differences)<={1,2,3,4,333},differences
 assert ISBN in pdf[2].get_text() and 'US$20.00' in pdf[2].get_text()
 assert all(1<=x[2]<=333 for x in pdf.get_toc())
 for i,p in enumerate(pdf):
  for l in p.get_links():
   if l.get('kind')==fitz.LINK_GOTO:assert 0<=l.get('page',-1)<333,(i,l)
 pdf.set_metadata({'title':TITLE+'——从封闭学生到人—智能器具复合主体','author':'王德生','subject':'德麦国际专著第289卷 | ISBN '+ISBN+' | US$20.00 | 校订v1.1','keywords':'SDE;教育发生学;复合主体;ISBN '+ISBN,'creator':'德麦国际出版社','producer':'LibreOffice / PyMuPDF'})
 pdf.set_page_labels([{'startpage':0,'prefix':'Cover'},{'startpage':1,'style':'D','firstpagenum':1},{'startpage':332,'prefix':'Back cover'}]);dest.mkdir(parents=True,exist_ok=True)
 printfile=dest/'education-subject-rebirth-print.pdf';printfile.unlink(missing_ok=True);pdf.save(printfile,garbage=4,deflate=True);pdf.close()
 pdf=fitz.open(printfile)
 for page in list(pdf)[1:-1]:page.draw_rect(page.rect,color=None,fill=(.9843,.9725,.9412),overlay=False)
 reader=dest/'education-subject-rebirth-reader.pdf';reader.unlink(missing_ok=True);pdf.save(reader,garbage=4,deflate=True);pdf.close();src.close()
 return {'physical_pages':333,'body_paper_pages':331,'bookmarks':240,'changed_text_pages':differences,'body_page_text_unchanged':True,'internal_links_valid':True}
def finalize_site(source,out,build,rendered):
 site_pages(source,out)
 for name in ['cover.jpg','backcover.jpg']:shutil.copy2(build/'assets'/name,out/name)
 shutil.copy2(build/'education-subject-rebirth.docx',out/'downloads/education-subject-rebirth.docx')
 pdf_check=finalize_pdf(rendered,source/'downloads/education-subject-rebirth-print.pdf',out/'downloads',build/'assets')
 source_blocks=signature_html((source/'text/index.html').read_text());assert source_blocks==signature_html((out/'text/index.html').read_text())
 mf=json.loads((source/'publication-manifest.json').read_text());mf.update(number=289,isbn=ISBN,price=20,priceCurrency='USD',edition='2026年10月第1版 · 校订 v1.1',canonical_url='https://sdeuniverses.com'+NEW,previous_url='https://sdeuniverses.com'+OLD)
 mf['files']={str(p.relative_to(out)):{'sha256':digest(p.read_bytes()),'bytes':p.stat().st_size} for p in sorted(out.rglob('*')) if p.is_file() and p.name not in ['publication-manifest.json','release-check.json']}
 mf['release_sha256']=digest(json.dumps(mf['files'],sort_keys=True).encode());(out/'publication-manifest.json').write_text(json.dumps(mf,ensure_ascii=False,indent=2))
 audit={'edition':'v1.1','volume':289,'isbn':ISBN,'isbn_checksum_valid':True,'price_usd':20,'source_markup_elements':len(source_blocks),'source_paragraphs_preserved':True,'files':len(mf['files']),**pdf_check};(out/'release-check.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2));return audit
