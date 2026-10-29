from pathlib import Path
import re,json,sys
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'output';OUT.mkdir(exist_ok=True)
net_han=0
for manuscript_file in (ROOT/'work/manuscript').glob('*.md'):
 content=re.sub(r'(?m)^\s*#+[^\n]*','',manuscript_file.read_text())
 content=re.sub(r'\[资料[：:].*?\]|（资料[：:].*?）','',content,flags=re.S)
 net_han+=len(re.findall(r'[\u4e00-\u9fff]',content))
doc=Document();sec=doc.sections[0]
sec.page_width=Mm(190);sec.page_height=Mm(250);sec.top_margin=Mm(18);sec.bottom_margin=Mm(18);sec.left_margin=Mm(20);sec.right_margin=Mm(18)
sec.header_distance=Mm(8);sec.footer_distance=Mm(8);sec.different_first_page_header_footer=True
styles=doc.styles
for st in styles:
 for el in list(st._element.iter(qn('w:pBdr'))):el.getparent().remove(el)
 for rf in st._element.iter(qn('w:rFonts')):
  for a in list(rf.attrib):
   if 'Theme' in a or 'theme' in a:del rf.attrib[a]
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=styles[name];st.font.name='Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor.from_string('242B32')
 for kind in ['ascii','hAnsi','cs']:st._element.get_or_add_rPr().rFonts.set(qn('w:'+kind),'Noto Serif CJK SC')
normal=styles['Normal'];normal.font.size=Pt(10.8);pf=normal.paragraph_format;pf.line_spacing=1.55;pf.space_after=Pt(5);pf.first_line_indent=Pt(21.6);pf.widow_control=True
for name,size in [('Title',27),('Heading 1',20),('Heading 2',14),('Heading 3',11.5)]:
 st=styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.space_before=Pt(18);st.paragraph_format.space_after=Pt(12);st.paragraph_format.keep_with_next=True
styles['Title'].font.color.rgb=RGBColor(0,0,0)
styles['Heading 1'].paragraph_format.page_break_before=True
for name in ['Heading 1','Heading 2']:styles[name].font.color.rgb=RGBColor.from_string('1F3A5F')
h=sec.header.paragraphs[0];h.text='创造力发生机制解码';h.style='Normal';h.paragraph_format.first_line_indent=Pt(0);h.paragraph_format.space_after=Pt(0)
for r in h.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('7C7A73')
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
for _ in range(3):doc.add_paragraph('')
p=doc.add_paragraph('德麦国际专著 第31卷');p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('创造力发生机制解码','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Decoding the Generative Mechanisms of Creativity','Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
p=doc.add_paragraph('王德生、Yang Bide 著');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
for _ in range(2):doc.add_paragraph('')
p=doc.add_paragraph('重写扩容版 2026年10月\n德麦国际有限公司\nDemai International Pte. Ltd.');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
doc.add_page_break();p=doc.add_paragraph('出版信息','Heading 2')
for txt in ['创造力发生机制解码','作者 王德生、Yang Bide','专著卷号 31','ISBN 978-1-970820-20-1','出版方 德麦国际有限公司 Demai International Pte. Ltd.','版本 2026年10月重写扩容版',f'正文 {net_han:,} 汉字（不含章节标题、目录、资料标注及参考资料）','本书由《三律方法学初稿》重写发展而来。理论论证、构造案例与拟议研究在正文中分别标明。','© 王德生、Yang Bide。保留相关权利。']:
 p=doc.add_paragraph(txt);p.paragraph_format.first_line_indent=Pt(0)
doc.add_page_break();doc.add_paragraph('目录','Heading 2')
chap=json.loads((ROOT/'work/chapter-map.json').read_text())
part_names=['三律哲学与发生学基础','目标路径约束与方法','创造的心理与组织机制','科学想象与人机共创','实践检验与知识谱系']
page_map_file=ROOT/'work/page-map.json'
page_map=json.loads(page_map_file.read_text()) if page_map_file.exists() else {}
citation_path=ROOT/'work/citation-map.json'
citation_map=json.loads(citation_path.read_text()) if citation_path.exists() else {}
items=[('intro','导论 创造力的发生问题')]+[(f'ch{x["number"]:02}',f'第{x["number"]}章 {x["title"]}') for x in chap]+[('conclusion','结语 让创造继续发生')]
if (ROOT/'work/glossary.json').exists():items.append(('glossary','附录 术语提要与阅读索引'))
if (ROOT/'work/references.md').exists():items.append(('references','参考资料与来源说明'))
for key,title in items:
 if key in ['ch01','ch13','ch25','ch37','ch49']:
  part=(int(key[2:])-1)//12
  p=doc.add_paragraph('第'+'一二三四五'[part]+'编 '+part_names[part]);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_with_next=True
  for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor.from_string('1F3A5F')
 p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.15
 p.paragraph_format.tab_stops.add_tab_stop(Mm(150),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),key);run=OxmlElement('w:r');t=OxmlElement('w:t');t.text=title;run.append(t);link.append(run);p._p.append(link)
 if key in page_map:p.add_run('\t'+str(page_map[key]))

def add_text(p,text):
 for i,v in enumerate(re.split(r'(\*\*.*?\*\*)',text)):
  r=p.add_run(v[2:-2] if v.startswith('**') and v.endswith('**') else v)
  if v.startswith('**') and v.endswith('**'):r.bold=True

def bookmark(p,key,num):
 a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(num));a.set(qn('w:name'),key);b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(num));p._p.insert(0,a);p._p.append(b)

files=[ROOT/'work/manuscript/introduction.md']+sorted((ROOT/'work/manuscript').glob('ch[0-9][0-9].md'))
if (ROOT/'work/manuscript/conclusion.md').exists():files.append(ROOT/'work/manuscript/conclusion.md')
for idx,path in enumerate(files,1):
 key='intro' if path.stem=='introduction' else path.stem
 text=path.read_text()
 for label,citation in citation_map.get(path.stem,{}).items():text=text.replace(label,citation)
 blocks=re.split(r'\n\s*\n',text.strip());section_no=0
 for block in blocks:
  block=block.strip()
  if not block:continue
  if block.startswith('# '):
   title=block[2:].strip()
   if key.startswith('ch'):
    number=int(key[2:]);title=f'第{number}章 '+chap[number-1]['title']
    if number in [1,13,25,37,49]:
     part=(number-1)//12
     p=doc.add_paragraph('第'+'一二三四五'[part]+'编 · '+part_names[part]);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
     for r in p.runs:r.font.size=Pt(11);r.font.color.rgb=RGBColor.from_string('9A7A3C')
   p=doc.add_paragraph(title,'Heading 1');bookmark(p,key,idx)
   if key in ['ch01','ch13','ch25','ch37','ch49']:p.paragraph_format.page_break_before=False
  elif block.startswith('## '):
   section_no+=1
   section_title=re.sub(r'^(?:第[一二三四1234]+节[、.．：:\s]*|[一二三四1234]+[、.．\s]+)','',block[3:].strip())
   doc.add_paragraph('第'+'一二三四'[section_no-1]+'节 '+section_title,'Heading 2')
  elif block.startswith('### '):doc.add_paragraph(block[4:].strip(),'Heading 3')
  elif block.startswith('```'):raise ValueError('Raw code block in manuscript '+path.name)
  elif block.startswith('|'):raise ValueError('Table needs native rendering '+path.name)
  else:
   p=doc.add_paragraph();add_text(p,re.sub(r'\n(?!\n)','',block))
refs=ROOT/'work/references.md'
glossary=ROOT/'work/glossary.json'
if glossary.exists():
 p=doc.add_paragraph('附录 术语提要与阅读索引','Heading 1');bookmark(p,'glossary',1001)
 p=doc.add_paragraph('下列提要用于回查概念及相关章节，不替代正文中的论证和条件。附录不计入净正文汉字数。');p.paragraph_format.first_line_indent=Pt(0)
 for entry in json.loads(glossary.read_text()):
  doc.add_paragraph(entry['term'],'Heading 3')
  p=doc.add_paragraph(entry['definition']+' 参见第'+'、'.join(map(str,entry['chapters']))+'章。');p.paragraph_format.first_line_indent=Pt(0)
if refs.exists():
 p=doc.add_paragraph('参考资料与来源说明','Heading 1');bookmark(p,'references',1002)
 for block in re.split(r'\n\s*\n',refs.read_text().strip()):
  if block.startswith('#'):continue
  p=doc.add_paragraph(block);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.3;p.paragraph_format.keep_together=True
props=doc.core_properties;props.title='创造力发生机制解码';props.author='王德生、Yang Bide';props.subject='SDE创造力发生学';props.keywords='创造力,特征律,自由律,幸福律,SDE';props.comments=''
file=OUT/'创造力发生机制解码_第31卷_正式文稿.docx';doc.save(file)
print(json.dumps({'file':str(file),'sections':len(files),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
